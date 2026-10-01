from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from amazon_agro.config.settings import AppSettings
from amazon_agro.domain.models import (
    Participant, Proposal, ProposalProperty, PropertyClassification,
)
from amazon_agro.exporters.excel import OpenpyxlExcelProposalExporter
from amazon_agro.exporters.pdf import SpreadsheetPdfProposalExporter
from amazon_agro.exporters.pdf_backends import (
    PdfBackendSelector, PdfBackendUnavailableError,
)
from amazon_agro.repositories.fake_properties import FakePropertyRepository
from amazon_agro.repositories.sqlite_proposals import SQLiteProposalRepository
from amazon_agro.services.export_service import ProposalExportService
from amazon_agro.services.export_validator import ProposalExportValidator
from amazon_agro.services.proposal_service import ProposalService


class FakeBackend:
    def __init__(self, name: str, available: bool = True, fail: bool = False) -> None:
        self.name = name
        self.available = available
        self.fail = fail
        self.proposal_number: str | None = None
        self.print_area: str | None = None
        self.footer: tuple[str | None, str | None] | None = None
        self.one_page_setup: bool | None = None

    def is_available(self) -> bool:
        return self.available

    def convert(self, filled_xlsx, destination):
        sheet = load_workbook(filled_xlsx).active
        self.proposal_number = sheet["K3"].value
        self.print_area = str(sheet.print_area)
        self.footer = (sheet["A34"].value, sheet["A35"].value)
        self.one_page_setup = (
            sheet.sheet_properties.pageSetUpPr.fitToPage is True
            and sheet.page_setup.fitToWidth == 1
            and sheet.page_setup.fitToHeight == 1
        )
        if self.fail:
            raise RuntimeError("Conversor indisponível")
        destination.write_bytes(b"%PDF-1.4\n% fake backend for tests\n")
        return destination


def _proposal() -> Proposal:
    proposal = Proposal(
        numero_proposta="123/45", proponente="João da Silva",
        cpf_cnpj="123.456.789-00", tecnico="Maria Técnica",
        finalidade="Custeio", cidade="Palmas", data_proposta=date(2026, 9, 29),
        valor_total=Decimal("1000.00"),
    )
    proposal.participants.append(Participant(proposal.id, "João da Silva"))
    proposal.properties.append(ProposalProperty(
        proposal.id, "DEMO-001", PropertyClassification.CLASS_1,
    ))
    return proposal


def test_pdf_backend_selection_keeps_preferred_order() -> None:
    unavailable = FakeBackend("unavailable", available=False)
    excel = FakeBackend("Excel COM")
    libreoffice = FakeBackend("LibreOffice")
    selector = PdfBackendSelector([unavailable, excel, libreoffice])
    assert selector.available_backends() == [excel, libreoffice]


def test_pdf_falls_back_and_uses_print_ready_copy(tmp_path) -> None:
    settings = AppSettings()
    xlsx = OpenpyxlExcelProposalExporter(
        settings, FakePropertyRepository()
    ).export(_proposal(), tmp_path / "proposal.xlsx")
    original_print_area = load_workbook(xlsx).active.print_area
    first = FakeBackend("Excel COM", fail=True)
    second = FakeBackend("LibreOffice")
    target = tmp_path / "proposal.pdf"
    SpreadsheetPdfProposalExporter(
        PdfBackendSelector([first, second])
    ).export(xlsx, target)
    assert target.read_bytes().startswith(b"%PDF")
    assert first.proposal_number == second.proposal_number == "123/45"
    assert "$A$1:$M$35" in second.print_area
    assert second.footer == (
        "TÉCNICO RESPONSÁVEL: Maria Técnica",
        "Palmas, 29 de setembro de 2026",
    )
    assert second.one_page_setup is True
    assert load_workbook(xlsx).active.print_area == original_print_area


def test_pdf_reports_when_no_converter_is_available(tmp_path) -> None:
    xlsx = OpenpyxlExcelProposalExporter(
        AppSettings(), FakePropertyRepository()
    ).export(_proposal(), tmp_path / "proposal.xlsx")
    exporter = SpreadsheetPdfProposalExporter(PdfBackendSelector([]))
    with pytest.raises(PdfBackendUnavailableError, match="Nenhum conversor PDF"):
        exporter.export(xlsx, tmp_path / "proposal.pdf")


def test_generate_both_uses_one_saved_proposal_and_filename(tmp_path) -> None:
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    properties = FakePropertyRepository()
    proposal = _proposal()
    ProposalService(repository, properties).save(proposal)
    backend = FakeBackend("fake")
    settings = AppSettings()
    service = ProposalExportService(
        repository, ProposalExportValidator(settings, properties),
        OpenpyxlExcelProposalExporter(settings, properties),
        SpreadsheetPdfProposalExporter(PdfBackendSelector([backend])),
    )
    result = service.generate_both(proposal.id, tmp_path)
    assert result.xlsx_path is not None
    assert result.pdf_path is not None
    assert result.xlsx_path.name == "123_45_Joao_da_Silva_proposta.xlsx"
    assert result.pdf_path.name == "123_45_Joao_da_Silva_proposta.pdf"
    assert load_workbook(result.xlsx_path).active["K3"].value == "123/45"
    assert backend.proposal_number == "123/45"
    assert result.pdf_path.read_bytes().startswith(b"%PDF")
    with pytest.raises(FileExistsError):
        service.generate_both(proposal.id, tmp_path)


def test_generate_both_publishes_nothing_if_pdf_fails(tmp_path) -> None:
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    properties = FakePropertyRepository()
    proposal = _proposal()
    ProposalService(repository, properties).save(proposal)
    settings = AppSettings()
    service = ProposalExportService(
        repository, ProposalExportValidator(settings, properties),
        OpenpyxlExcelProposalExporter(settings, properties),
        SpreadsheetPdfProposalExporter(PdfBackendSelector([
            FakeBackend("failing", fail=True)
        ])),
    )
    with pytest.raises(RuntimeError, match="Não foi possível gerar o PDF"):
        service.generate_both(proposal.id, tmp_path)
    assert not list(tmp_path.glob("*_proposta.xlsx"))
    assert not list(tmp_path.glob("*_proposta.pdf"))


@pytest.mark.parametrize("fail", [False, True])
def test_excel_releases_com_interfaces_before_uninitializing(tmp_path, monkeypatch, fail):
    import sys
    from types import ModuleType
    from amazon_agro.exporters.pdf_backends import ExcelComPdfBackend
    events = []
    class Book:
        def ExportAsFixedFormat(self, *args):
            if fail:
                raise RuntimeError("conversion error")
        def Close(self, **kwargs):
            assert kwargs == {"SaveChanges": False}
            events.append("close")
        def __del__(self):
            events.append("release-book")
    class Application:
        @property
        def Workbooks(self): return self
        def Open(self, *args, **kwargs): return Book()
        def Quit(self): events.append("quit")
        def __del__(self): events.append("release-application")
    pythoncom = ModuleType("pythoncom")
    pythoncom.CoInitialize = lambda: events.append("initialize")
    def uninitialize():
        # An exception traceback may retain the failing proxy until propagated.
        assert "close" in events and "quit" in events
        if not fail:
            assert "release-book" in events
            assert "release-application" in events
        events.append("uninitialize")
    pythoncom.CoUninitialize = uninitialize
    client = ModuleType("win32com.client")
    client.DispatchEx = lambda name: Application()
    monkeypatch.setitem(sys.modules, "pythoncom", pythoncom)
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    monkeypatch.setattr(ExcelComPdfBackend, "is_available", lambda self: True)
    backend = ExcelComPdfBackend()
    if fail:
        with pytest.raises(RuntimeError, match="conversion error"):
            backend.convert(tmp_path / "source.xlsx", tmp_path / "result.pdf")
    else:
        backend.convert(tmp_path / "source.xlsx", tmp_path / "result.pdf")
    assert "uninitialize" in events
    assert events.index("uninitialize") > events.index("close")
    assert events.index("uninitialize") > events.index("quit")
