"""Opt-in diagnostics through the real application and services, synthetic data only."""
from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from decimal import Decimal
from hashlib import sha256
import io
import json
import logging
import os
from pathlib import Path
import sqlite3
import subprocess
import sys
import time

from PySide6.QtCore import Qt, QTimer
from PySide6.QtTest import QTest, QSignalSpy
from PySide6.QtWidgets import QApplication, QFileDialog, QLineEdit, QMessageBox, QWizard
from openpyxl import Workbook, load_workbook

from amazon_agro.config.settings import AppSettings, default_data_dir, default_config_path
from amazon_agro.ui.export_dialogs import ExportSuccessDialog
from amazon_agro.ui.onboarding import FirstRunWizard
from amazon_agro.version import __version__

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class SmokeOptions:
    directory: Path
    pdf: bool = False
    no_pdf: bool = False


def configure_smoke(arguments: list[str]) -> SmokeOptions | None:
    if "--smoke-test" not in arguments:
        return None
    parser = argparse.ArgumentParser(description="Diagnóstico isolado com dados fictícios")
    parser.add_argument("--smoke-test", required=True, type=Path)
    parser.add_argument("--smoke-pdf", action="store_true")
    parser.add_argument("--smoke-no-pdf", action="store_true")
    options = parser.parse_args(arguments)
    directory = options.smoke_test.resolve()
    directory.mkdir(parents=True, exist_ok=False)
    # Process-local overrides; never touch the operator's real settings or data.
    os.environ["LOCALAPPDATA"] = str(directory / "localappdata")
    os.environ["AMAZON_AGRO_CONFIG"] = str(directory / "localappdata/AmazonAgro/config.json")
    return SmokeOptions(directory, options.smoke_pdf, options.smoke_no_pdf)


def excel_process_ids() -> set[int]:
    tasklist = Path(os.environ.get("SystemRoot", "C:/Windows")) / "System32/tasklist.exe"
    if sys.platform != "win32" or not tasklist.is_file():
        return set()
    result = subprocess.run([str(tasklist), "/FI", "IMAGENAME eq EXCEL.EXE", "/FO", "CSV", "/NH"],
                            capture_output=True, text=True, check=True,
                            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    return {int(row[1]) for row in csv.reader(io.StringIO(result.stdout))
            if len(row) >= 2 and row[0].lower() == "excel.exe"}


class SmokeDriver:
    def __init__(self, window, options: SmokeOptions, app: QApplication):
        self.window, self.options, self.app = window, options, app
        self.destination = options.directory / "xlsx"
        self.paths: list[Path] = []
        self.dialog_errors: list[str] = []
        self.unavailable_seen = False
        self.timer = QTimer(window)
        self.timer.setInterval(80)
        self.timer.timeout.connect(self.handle_dialog)

    def handle_dialog(self) -> None:
        dialog = self.app.activeModalWidget()
        if isinstance(dialog, FirstRunWizard):
            index = dialog.currentId()
            if index == 1:
                dialog.source.setText(str(self.options.directory / "source"))
            elif index == 2:
                dialog.output.setText(str(self.options.directory / "exports"))
            elif index == 4:
                dialog.city.setText("Palmas")
                dialog.technician.setText("Tecnico Exemplo")
            button = QWizard.WizardButton.FinishButton if index == 5 else QWizard.WizardButton.NextButton
            QTest.mouseClick(dialog.button(button), Qt.MouseButton.LeftButton)
        elif isinstance(dialog, QFileDialog):
            dialog.setDirectory(str(self.destination))
            dialog.findChild(QLineEdit, "fileNameEdit").setText(str(self.destination))
            dialog.accept()
        elif isinstance(dialog, ExportSuccessDialog):
            self.paths.extend(dialog.paths)
            dialog.accept()
        elif isinstance(dialog, QMessageBox):
            if dialog.windowTitle() == "PDF indisponível":
                self.unavailable_seen = True
                close = next(button for button in dialog.buttons() if button.text() == "Fechar")
                QTest.mouseClick(close, Qt.MouseButton.LeftButton)
            else:
                self.dialog_errors.append(dialog.text())
                dialog.reject()

    def export(self, kind: str) -> list[Path]:
        self.destination = self.options.directory / kind
        self.destination.mkdir(exist_ok=True)
        self.paths.clear()
        self.window.settings.default_output_dir = str(self.destination)
        self.window.steps.setCurrentRow(4)
        self.app.processEvents()
        button = {"xlsx": self.window.review_page.excel_button,
                  "pdf": self.window.review_page.pdf_button,
                  "both": self.window.review_page.both_button}[kind]
        spy = QSignalSpy(button.clicked)
        QTest.mouseClick(button, Qt.MouseButton.LeftButton)
        if spy.count() != 1 or self.dialog_errors:
            raise RuntimeError(f"UI export did not complete: {self.dialog_errors}")
        return list(self.paths)

    def run(self) -> None:
        from sqlalchemy.util._has_cython import HAS_CYEXTENSION
        root = self.options.directory
        report = {"ok": False, "frozen": bool(getattr(sys, "frozen", False)),
                  "version": __version__, "platform": self.app.platformName(),
                  "executable": sys.executable, "cwd": str(Path.cwd()),
                  "data_dir": str(default_data_dir()), "config_path": str(default_config_path()),
                  "sqlalchemy_pure_python": not HAS_CYEXTENSION,
                  "external_python_on_path": any("python" in part.lower() for part in os.environ.get("PATH", "").split(os.pathsep))}
        try:
            source = root / "source"
            source.mkdir()
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "BASA Ambiental"
            sheet.append(["Fazenda", "Área", "Matrículas", "Matrícula Anterior", "Lote/Gleba", "Proprietário", "CPF/CNPJ", "Município", "UF"])
            sheet.append(["Fazenda Exemplo", "10,00", "MAT-005", "", "Lote Exemplo", "Pessoa Exemplo", "000.000.000-00", "Palmas", "TO"])
            workbook.save(source / "synthetic.xlsx")
            workbook.close()
            if self.options.no_pdf:
                self.window.export_service.pdf.selector.backends = []
            self.timer.start()
            self.window.startup()
            settings = self.window.settings
            if not settings.first_run_completed or not self.window.has_proposal:
                raise RuntimeError("Fresh-install onboarding did not complete.")
            loaded_settings = AppSettings.load()
            if loaded_settings.source_directory() != source or not loaded_settings.first_run_completed:
                raise RuntimeError("Onboarding preferences were not persisted.")
            self.window.startup()  # Completed configuration must not show the wizard again.
            report["onboarding_ok"] = True
            self.window.sync_service.synchronize(settings)
            self.window.catalog.search_enabled = True
            self.window.properties_page.search()
            if self.window.properties_page.results.rowCount() != 1:
                raise RuntimeError("Local synchronized-folder source did not load.")
            self.window.properties_page.results.setCurrentCell(0, 0)
            farm = self.window.properties_page._current_property()
            if not farm.municipality or not farm.state:
                self.window.service.save_property_local_enrichment(farm.external_id, "Palmas", "TO")
                farm = self.window.service.get_property(farm.external_id)
            from amazon_agro.domain.models import ProposalProperty, ProposalPropertyParcel, PropertyClassification
            self.window.properties_page._add_link(ProposalProperty(
                self.window.current.id, farm.external_id, PropertyClassification.CREDIT_OBJECT,
                property_name_snapshot=farm.name, municipality_snapshot=farm.municipality,
                state_snapshot=farm.state, source_file_snapshot=farm.source_file,
                selected_parcels=[ProposalPropertyParcel(farm.parcels[0].external_id, farm.parcels[0].registration)],
            ))
            from amazon_agro.domain.models import ParticipantType
            kind = self.window.participants_page.table.cellWidget(0, 2)
            assert kind.currentData() == int(ParticipantType.MAIN_ISSUER)
            assert self.window.participants_page.table.item(0, 1).text() == "07.778.284/0001-90"
            report["local_source_ok"] = True
            for name, value in {"numero_proposta": "TEST-005", "proponente": "Proponente Exemplo",
                                "cpf_cnpj": "000.000.000-00", "finalidade": "Custeio", "agencia": "Agencia Exemplo"}.items():
                control = self.window.operation_page.controls[name]
                (control.setCurrentText if hasattr(control, "setCurrentText") else control.setText)(value)
            self.window.proposal_page.spins["valor_total"].setValue(Decimal("100000.00"))
            self.window.proposal_page.spins["valor_fno"].setValue(Decimal("70000.00"))
            self.window.proposal_page.spins["valor_of"].setValue(Decimal("30000.00"))
            self.window.proposal_page.abc_yes.setChecked(True)
            self.window.proposal_page.spins["laudo_abc_percentual"].setValue(5)
            assert self.window.proposal_page.abc_amount.text() == "R$ 3.500,00"
            self.window.proposal_page.spins["valor_fno"].setValue(Decimal("80000"))
            assert self.window.proposal_page.abc_amount.text() == "R$ 4.000,00"
            self.window.proposal_page.spins["valor_fno"].setValue(Decimal("70000"))
            self.window.proposal_page.spins["valor_total"].setValue(Decimal("120000"))
            assert self.window.proposal_page.abc_amount.text() == "R$ 3.500,00"
            self.window.proposal_page.spins["valor_total"].setValue(Decimal("100000"))
            for source, percent in (("fno", 2.75), ("of", 1.25)):
                flag = f"astec_{source}_financiada"
                fields = self.window.proposal_page.astec_fields[flag]
                assert fields.isHidden() and not fields.isEnabled()
                self.window.proposal_page.flags[flag].setValue(True)
                self.window.proposal_page.spins[f"astec_{source}_percentual"].setValue(percent)
                assert not fields.isHidden() and fields.isEnabled()
            report["abc_fno_basis_ok"] = report["abc_total_independent_ok"] = True
            assert "laudo_abc_valor" not in self.window.proposal_page.spins
            assert "recursos_proprios" not in self.window.proposal_page.spins
            report["abc_calculation_ok"] = True
            self.window.save_proposal()
            saved = self.window.service.get(self.window.current.id)
            if saved is None:
                raise RuntimeError("Proposal was not persisted.")
            assert len(saved.participants) == 1
            assert saved.participants[0].tipo == ParticipantType.MAIN_ISSUER
            assert saved.laudo_abc_valor == Decimal("3500.00")
            assert saved.laudo_abc_percentual == Decimal("5")
            self.window._load(saved)
            assert saved.astec_fno_financiada and saved.astec_of_financiada
            assert saved.astec_fno_percentual == Decimal("2.75")
            assert saved.astec_of_percentual == Decimal("1.25")
            for fields in self.window.proposal_page.astec_fields.values():
                assert not fields.isHidden() and fields.isEnabled()
            report["astec_persistence_ok"] = True
            report["save_reopen_ok"] = True
            report["abc_persistence_ok"] = True
            template = settings.template_path()
            before_hash = sha256(template.read_bytes()).hexdigest()
            report["template"] = str(template)
            xlsx_paths = self.export("xlsx")
            if len(xlsx_paths) != 1 or not xlsx_paths[0].is_file():
                raise RuntimeError("No final XLSX was produced.")
            output = load_workbook(xlsx_paths[0])
            sheet = output.active
            assert sheet["A3"].value == "TEST-005" and sheet["A6"].value == "Proponente Exemplo"
            assert sheet["A28"].value == "Fazenda Exemplo"
            assert sheet["A9"].value == settings.default_consultancy_name
            assert sheet["I9"].value == "07.778.284/0001-90" and sheet["M9"].value == 1
            assert sheet["I23"].value == "Sim" and sheet["K23"].value == 0.05
            assert sheet["I25"].value == 3500
            assert sheet["F23"].value == sheet["F25"].value == "Sim"
            assert sheet["H23"].value == 0.0275 and sheet["H25"].value == 0.0125
            assert "SOMENTE FNO" in sheet["I24"].value
            report["astec_export_ok"] = True
            assert sheet["I22"].value == "LAUDO ABC FINANCIADO?"
            for address, percentage in (("D23", saved.fno_percentage), ("D25", saved.of_percentage)):
                assert abs(Decimal(str(sheet[address].value)) - percentage / Decimal("100")) < Decimal("0.00000000000001")
            assert len(sheet._images) == 1 and sheet["A2"].value != "LOGO AMAZON"
            assert sheet.page_setup.fitToWidth == sheet.page_setup.fitToHeight == 1
            assert sheet.sheet_properties.pageSetUpPr.fitToPage
            output.close()
            report["xlsx_ok"] = True
            report["amazon_default_ok"] = True
            report["abc_export_ok"] = True
            report["financing_shares_ok"] = True
            report["official_logo_ok"] = True
            report["amazon_default_type"] = 1
            report["xlsx_path"] = str(xlsx_paths[0])
            available = self.window.export_service.pdf.selector.available_backends()
            report["pdf_backends"] = [backend.name for backend in available]
            if self.options.no_pdf or (self.options.pdf and not available):
                self.export("pdf")
                assert self.unavailable_seen
                report["pdf_unavailable_feedback"] = True
            elif self.options.pdf:
                initial_excel = excel_process_ids()
                report["pdf_files"] = []
                for kind in ("pdf", "both"):
                    paths = self.export(kind)
                    assert len(paths) == (1 if kind == "pdf" else 2)
                    for path in paths:
                        assert path.is_file() and path.stat().st_size > 0
                        if path.suffix == ".pdf":
                            assert path.read_bytes().startswith(b"%PDF")
                            report["pdf_files"].append(str(path))
                deadline = time.monotonic() + 10
                remaining = excel_process_ids() - initial_excel
                while remaining and time.monotonic() < deadline:
                    self.app.processEvents()
                    time.sleep(0.2)
                    remaining = excel_process_ids() - initial_excel
                report["new_excel_processes_remaining"] = sorted(remaining)
                if remaining:
                    raise RuntimeError("Excel left a new background process; inspect the conversion log.")
                report["pdf_ok"] = True
            assert sha256(template.read_bytes()).hexdigest() == before_hash
            report["template_unchanged"] = True
            for name in ("proposals.sqlite3", "properties.sqlite3"):
                with sqlite3.connect(default_data_dir() / name) as connection:
                    assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
                    assert connection.execute("PRAGMA integrity_check").fetchone() == ("ok",)
            report["sqlite_integrity"] = "ok"
            report["ok"] = True
        except Exception as error:
            logger.exception("Packaged smoke failed")
            report["error"] = f"{type(error).__name__}: {error}"
        finally:
            self.timer.stop()
            (root / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            self.app.exit(0 if report["ok"] else 1)
