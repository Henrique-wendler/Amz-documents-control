from __future__ import annotations

from hashlib import sha256
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from amazon_agro.config.settings import AppSettings
from amazon_agro.integrations.workbook_discovery import PropertyWorkbookDiscovery
from amazon_agro.integrations.workbook_inspector import (
    NeedsConfigurationError, WorkbookInspector,
)
from amazon_agro.integrations.workbook_parser import PropertyWorkbookParser
from amazon_agro.integrations.workbook_profile import PropertyWorkbookProfile
from amazon_agro.repositories.sqlite_property_catalog import SQLitePropertyCatalogRepository
from amazon_agro.services.property_sync_service import PropertyCatalogSyncService
from amazon_agro.ui.privacy import mask_document


HEADERS = (
    "Fazenda", "Área (ha)", "Matrículas", "Matrícula Anterior", "Lote/Gleba",
    "Proprietário", "CPF/CNPJ", "CCIR", "ITR", "CAR", "Município",
)


def record(name: str = "Fazenda Exemplo", registration: str = "MAT-1", **values):
    result = {
        "Fazenda": name, "Área (ha)": 10.5, "Matrículas": registration,
        "Matrícula Anterior": "ANT-1", "Lote/Gleba": "Lote A",
        "Proprietário": "Pessoa Exemplo", "CPF/CNPJ": "000.000.000-00",
        "CCIR": "CCIR-1", "ITR": "ITR-1", "CAR": "CAR-1",
        "Município": "Palmas",
    }
    result.update(values)
    return result


def make_book(
    path: Path, records: list[dict], *, header_row: int = 3,
    first_column: int = 2, headers: tuple[str, ...] = HEADERS,
    merges: tuple[tuple[str, int, int], ...] = (),
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dados"
    for offset, header in enumerate(headers):
        sheet.cell(header_row, first_column + offset, header)
    for row_offset, values in enumerate(records, 1):
        for offset, header in enumerate(headers):
            if header in values:
                sheet.cell(header_row + row_offset, first_column + offset, values[header])
    for header, first, last in merges:
        column = first_column + headers.index(header)
        sheet.merge_cells(
            start_row=header_row + first, end_row=header_row + last,
            start_column=column, end_column=column,
        )
    workbook.save(path)
    return path


def parse(path: Path, profile: PropertyWorkbookProfile | None = None):
    return PropertyWorkbookParser().parse(
        path, path.name, profile or PropertyWorkbookProfile()
    )


def catalog_sync(tmp_path: Path):
    catalog = SQLitePropertyCatalogRepository(tmp_path / "catalog.sqlite3")
    service = PropertyCatalogSyncService(catalog)
    settings = AppSettings(property_source_directory=str(tmp_path))
    return catalog, service, settings


def test_one_farm_one_registration(tmp_path) -> None:
    path = make_book(tmp_path / "one.xlsx", [record()])
    before = sha256(path.read_bytes()).digest()
    parsed = parse(path)
    assert len(parsed.properties) == 1
    farm = parsed.properties[0]
    assert farm.name == "Fazenda Exemplo"
    assert farm.parcels[0].registration == "MAT-1"
    assert str(farm.total_area) == "10.5"
    assert sha256(path.read_bytes()).digest() == before


def test_one_farm_many_registrations_with_merged_name(tmp_path) -> None:
    path = make_book(tmp_path / "many.xlsx", [
        record(),
        record(name=None, registration="MAT-2", **{
            "Proprietário": None, "CPF/CNPJ": None,
            "Matrícula Anterior": "ANT-2", "Lote/Gleba": "Lote B",
        }),
        record(name=None, registration="MAT-3", **{
            "Proprietário": None, "CPF/CNPJ": None,
        }),
    ], merges=(("Fazenda", 1, 3), ("Proprietário", 1, 3), ("CPF/CNPJ", 1, 3)))
    parsed = parse(path)
    assert len(parsed.properties) == 1
    assert [item.registration for item in parsed.properties[0].parcels] == [
        "MAT-1", "MAT-2", "MAT-3"
    ]


def test_multiple_farms_share_file_owner_but_keep_distinct_ids(tmp_path) -> None:
    path = make_book(tmp_path / "farms.xlsx", [
        record(name="Fazenda A", registration="A-1"),
        record(name="Fazenda B", registration="B-1", **{
            "Proprietário": None, "CPF/CNPJ": None,
        }),
    ], merges=(("Proprietário", 1, 2), ("CPF/CNPJ", 1, 2)))
    parsed = parse(path)
    assert len(parsed.properties) == 2
    assert parsed.owner_name == "Pessoa Exemplo"
    assert len({item.external_id for item in parsed.properties}) == 2
    assert all(len(item.parcels) == 1 for item in parsed.properties)


def test_property_metadata_propagates_only_through_merge(tmp_path) -> None:
    path = make_book(tmp_path / "merged.xlsx", [
        record(),
        record(name=None, registration="MAT-2", **{
            "Proprietário": None, "CPF/CNPJ": None, "CCIR": None,
            "ITR": None, "CAR": None, "Matrícula Anterior": None,
            "Lote/Gleba": None,
        }),
    ], merges=tuple((field, 1, 2) for field in (
        "Fazenda", "Proprietário", "CPF/CNPJ", "CCIR", "ITR", "CAR"
    )))
    farm = parse(path).properties[0]
    assert farm.ccir == "CCIR-1" and farm.itr == "ITR-1" and farm.car == "CAR-1"
    assert farm.parcels[1].previous_registration == ""
    assert farm.parcels[1].lot_description == ""


def test_unmerged_blank_farm_is_not_forward_filled(tmp_path) -> None:
    path = make_book(tmp_path / "ambiguous.xlsx", [
        record(), record(name=None, registration="MAT-2"),
    ])
    with pytest.raises(NeedsConfigurationError, match="sem fazenda"):
        parse(path)


def test_unmerged_blank_owner_is_not_forward_filled(tmp_path) -> None:
    path = make_book(tmp_path / "owner-gap.xlsx", [
        record(), record(name=None, registration="MAT-2", **{
            "Proprietário": None, "CPF/CNPJ": None,
        }),
    ], merges=(("Fazenda", 1, 2),))
    farm = parse(path).properties[0]
    assert len(farm.parcels) == 2
    assert farm.owner_name == "Pessoa Exemplo"


def test_contiguous_grouping_can_be_enabled_by_profile(tmp_path) -> None:
    path = make_book(tmp_path / "configured.xlsx", [
        record(), record(name=None, registration="MAT-2"),
    ])
    profile = PropertyWorkbookProfile.configured(
        "BASA Ambiental", {"inherit_unmerged_property_name": True}
    )
    assert len(parse(path, profile).properties[0].parcels) == 2


def test_separate_repeated_farm_blocks_need_configuration(tmp_path) -> None:
    path = make_book(tmp_path / "duplicate-farm.xlsx", [
        record(name="Fazenda A", registration="A-1"),
        record(name="Fazenda B", registration="B-1"),
        record(name="Fazenda A", registration="A-2"),
    ])
    with pytest.raises(NeedsConfigurationError, match="Blocos separados"):
        parse(path)


def test_area_merged_across_registrations_is_not_guessed(tmp_path) -> None:
    path = make_book(tmp_path / "shared-area.xlsx", [
        record(), record(name=None, registration="MAT-2", **{
            "Área (ha)": None, "Proprietário": None, "CPF/CNPJ": None,
        }),
    ], merges=(
        ("Fazenda", 1, 2), ("Proprietário", 1, 2),
        ("CPF/CNPJ", 1, 2), ("Área (ha)", 1, 2),
    ))
    with pytest.raises(NeedsConfigurationError, match="Área mesclada"):
        parse(path)


def test_headers_can_move_and_aliases_can_be_configured(tmp_path) -> None:
    headers = tuple("Estabelecimento" if item == "Fazenda" else item for item in HEADERS)
    path = make_book(
        tmp_path / "shifted.xlsx", [{**record(), "Estabelecimento": "Fazenda Móvel"}],
        header_row=8, first_column=5, headers=headers,
    )
    profile = PropertyWorkbookProfile.configured("Família", {
        "column_aliases": {"name": ["Estabelecimento"]},
    })
    assert parse(path, profile).properties[0].name == "Fazenda Móvel"


def test_identity_does_not_depend_on_excel_row_number(tmp_path) -> None:
    path = make_book(tmp_path / "stable.xlsx", [record()], header_row=3)
    first = parse(path).properties[0]
    make_book(path, [record()], header_row=9)
    second = parse(path).properties[0]
    assert first.external_id == second.external_id
    assert first.parcels[0].external_id == second.parcels[0].external_id


def test_owner_document_punctuation_does_not_change_property_id(tmp_path) -> None:
    path = make_book(tmp_path / "document.xlsx", [record()])
    first = parse(path).properties[0].external_id
    make_book(path, [record(**{"CPF/CNPJ": "00000000000"})])
    assert parse(path).properties[0].external_id == first


def test_owner_can_be_extracted_from_configured_cell(tmp_path) -> None:
    headers = tuple(item for item in HEADERS if item != "Proprietário")
    path = make_book(tmp_path / "owner-cell.xlsx", [record()], headers=headers)
    workbook = load_workbook(path)
    workbook.active["B1"] = "Pessoa Exemplo"
    workbook.save(path)
    profile = PropertyWorkbookProfile.configured("Família", {
        "owner_name_cell": "B1",
    })
    assert parse(path, profile).properties[0].owner_name == "Pessoa Exemplo"


def test_incompatible_headers_need_configuration(tmp_path) -> None:
    headers = tuple(item for item in HEADERS if item != "Matrículas")
    path = make_book(tmp_path / "wrong.xlsx", [record()], headers=headers)
    with pytest.raises(NeedsConfigurationError, match="Cabeçalhos"):
        parse(path)


def test_multiple_compatible_sheets_require_sheet_selector(tmp_path) -> None:
    path = make_book(tmp_path / "two-sheets.xlsx", [record()])
    workbook = load_workbook(path)
    copy = workbook.copy_worksheet(workbook.active)
    copy.title = "Outra aba"
    workbook.save(path)
    with pytest.raises(NeedsConfigurationError, match="Mais de uma aba"):
        parse(path)
    profile = PropertyWorkbookProfile.configured("Família", {
        "sheet_selector": "Dados",
    })
    assert len(parse(path, profile).properties) == 1


def test_inspector_reports_sheets_headers_merges_and_masked_preview(tmp_path) -> None:
    path = make_book(tmp_path / "inspect.xlsx", [
        record(), record(name=None, registration="MAT-2"),
    ], merges=(("Fazenda", 1, 2),))
    result = WorkbookInspector().inspect(path, PropertyWorkbookProfile())
    sheet = result.sheets[0]
    assert sheet.name == "Dados"
    assert sheet.header_candidates[0].row == 3
    assert "B4:B5" in sheet.merged_ranges
    assert sheet.column_names[1] == "Fazenda"
    assert "000.000.000-00" not in str(sheet.first_rows)
    assert "Pessoa Exemplo" not in str(sheet.first_rows)


def test_discovery_ignores_temporary_hidden_and_excluded_files(tmp_path) -> None:
    make_book(tmp_path / "ok.xlsx", [record()])
    for name in ("~$lock.xlsx", ".hidden.xlsx", "skip.xlsx"):
        (tmp_path / name).write_bytes(b"temporary")
    files = PropertyWorkbookDiscovery().discover(
        tmp_path, ["*.xlsx"], excluded_files=["skip.xlsx"]
    )
    assert [item.name for item in files] == ["ok.xlsx"]
    assert files[0].size > 0 and files[0].modified_at.tzinfo is not None


def test_discovery_supports_xlsm_and_optional_recursion(tmp_path) -> None:
    child = tmp_path / "sub"
    child.mkdir()
    make_book(child / "nested.xlsm", [record()])
    discovery = PropertyWorkbookDiscovery()
    assert discovery.discover(tmp_path, ["*.xlsm"], recursive=False) == []
    assert discovery.discover(tmp_path, ["*.xlsm"], recursive=True)[0].relative_path == (
        "sub/nested.xlsm"
    )


def test_sync_reads_xlsm_without_changing_it(tmp_path) -> None:
    path = make_book(tmp_path / "macro-family.xlsm", [record()])
    original_hash = sha256(path.read_bytes()).digest()
    catalog, sync, settings = catalog_sync(tmp_path)
    report = sync.synchronize(settings)
    assert report.stats.properties == 1
    assert sha256(path.read_bytes()).digest() == original_hash
    catalog.close()


def test_sync_imports_multiple_files_then_skips_unchanged(tmp_path) -> None:
    make_book(tmp_path / "first.xlsx", [record(name="Fazenda A")])
    make_book(tmp_path / "second.xlsx", [record(name="Fazenda B")])
    catalog, sync, settings = catalog_sync(tmp_path)
    first = sync.synchronize(settings)
    second = sync.synchronize(settings)
    assert [event.action for event in first.events] == ["NEW", "NEW"]
    assert [event.action for event in second.events] == ["UNCHANGED", "UNCHANGED"]
    assert second.stats.properties == second.stats.parcels == 2
    catalog.close()


def test_different_file_families_use_configured_profiles(tmp_path) -> None:
    make_book(tmp_path / "normal.xlsx", [record(name="Fazenda Padrão")])
    special_headers = tuple(
        "Estabelecimento" if item == "Fazenda" else item for item in HEADERS
    )
    make_book(tmp_path / "especial.xlsx", [{
        **record(), "Estabelecimento": "Fazenda Especial"
    }], headers=special_headers)
    catalog, sync, settings = catalog_sync(tmp_path)
    settings.property_profiles = {
        "Outra Família": {"column_aliases": {"name": ["Estabelecimento"]}}
    }
    settings.property_profile_rules = [
        {"pattern": "especial*.xlsx", "profile": "Outra Família"}
    ]
    report = sync.synchronize(settings)
    assert report.stats.properties == 2
    assert all(event.status == "OK" for event in report.events)
    assert catalog.get_source(tmp_path / "especial.xlsx").profile == "Outra Família"
    catalog.close()


def test_ambiguous_profile_rules_mark_file_for_configuration(tmp_path) -> None:
    path = make_book(tmp_path / "ambiguous.xlsx", [record()])
    catalog, sync, settings = catalog_sync(tmp_path)
    settings.property_profiles = {"Outra Família": {}}
    settings.property_profile_rules = [
        {"pattern": "*.xlsx", "profile": "BASA Ambiental"},
        {"pattern": "ambiguous.xlsx", "profile": "Outra Família"},
    ]
    sync.synchronize(settings)
    assert catalog.get_source(path).status == "NEEDS_CONFIGURATION"
    assert catalog.search("") == []
    catalog.close()


def test_exact_file_profile_overrides_family_rule(tmp_path) -> None:
    headers = tuple("Estabelecimento" if item == "Fazenda" else item for item in HEADERS)
    path = make_book(tmp_path / "special.xlsx", [{
        **record(), "Estabelecimento": "Fazenda Especial"
    }], headers=headers)
    catalog, sync, settings = catalog_sync(tmp_path)
    settings.property_profiles = {
        "Outra Família": {"column_aliases": {"name": ["Estabelecimento"]}}
    }
    settings.property_profile_rules = [
        {"pattern": "*.xlsx", "profile": "BASA Ambiental"}
    ]
    settings.property_file_profiles = {"special.xlsx": "Outra Família"}
    report = sync.synchronize(settings)
    assert report.stats.properties == 1
    assert catalog.get_source(path).profile == "Outra Família"
    catalog.close()


def test_sync_reprocesses_only_modified_file(tmp_path) -> None:
    first = make_book(tmp_path / "first.xlsx", [record(registration="OLD")])
    make_book(tmp_path / "second.xlsx", [record(name="Fazenda B")])
    catalog, sync, settings = catalog_sync(tmp_path)
    sync.synchronize(settings)
    workbook = load_workbook(first)
    workbook.active["D4"] = "NEW"
    workbook.save(first)
    report = sync.synchronize(settings)
    assert [(event.file, event.action) for event in report.events] == [
        ("first.xlsx", "MODIFIED"), ("second.xlsx", "UNCHANGED")
    ]
    assert len(catalog.search("NEW")) == 1
    assert not catalog.search("OLD")
    catalog.close()


def test_profile_change_reprocesses_unchanged_workbook(tmp_path) -> None:
    make_book(tmp_path / "profile.xlsx", [record()])
    catalog, sync, settings = catalog_sync(tmp_path)
    sync.synchronize(settings)
    settings.property_profile_options = {"header_scan_rows": 12}
    report = sync.synchronize(settings)
    assert [event.action for event in report.events] == ["MODIFIED"]
    assert report.stats.properties == 1
    catalog.close()


def test_removed_file_becomes_inactive_without_deleting_catalog_record(tmp_path) -> None:
    path = make_book(tmp_path / "removed.xlsx", [record()])
    catalog, sync, settings = catalog_sync(tmp_path)
    sync.synchronize(settings)
    property_id = catalog.search("")[0].external_id
    path.unlink()
    report = sync.synchronize(settings)
    assert report.events[0].action == "REMOVED"
    assert catalog.search("") == []
    assert catalog.get_by_external_id(property_id) is not None
    assert catalog.get_source(path).status == "REMOVED"
    catalog.close()


def test_corrupt_workbook_is_error_and_does_not_block_good_file(tmp_path) -> None:
    (tmp_path / "bad.xlsx").write_bytes(b"not an Excel zip")
    make_book(tmp_path / "good.xlsx", [record()])
    catalog, sync, settings = catalog_sync(tmp_path)
    report = sync.synchronize(settings)
    assert [(event.file, event.status) for event in report.events] == [
        ("bad.xlsx", "ERROR"), ("good.xlsx", "OK")
    ]
    assert report.stats.properties == 1
    catalog.close()


def test_incompatible_file_is_needs_configuration(tmp_path) -> None:
    headers = tuple(item for item in HEADERS if item != "Matrículas")
    path = make_book(tmp_path / "incompatible.xlsx", [record()], headers=headers)
    catalog, sync, settings = catalog_sync(tmp_path)
    sync.synchronize(settings)
    assert catalog.get_source(path).status == "NEEDS_CONFIGURATION"
    assert catalog.search("") == []
    catalog.close()


def test_failed_update_keeps_last_valid_catalog(tmp_path) -> None:
    path = make_book(tmp_path / "changing.xlsx", [record()])
    catalog, sync, settings = catalog_sync(tmp_path)
    sync.synchronize(settings)
    property_id = catalog.search("")[0].external_id
    path.write_bytes(b"corrupt update")
    sync.synchronize(settings)
    assert catalog.get_source(path).status == "ERROR"
    assert catalog.search("")[0].external_id == property_id
    assert catalog.search("")[0].source_status == "ERROR"
    catalog.close()


def test_similar_filenames_and_identical_farms_have_distinct_ids(tmp_path) -> None:
    make_book(tmp_path / "farm.xlsx", [record()])
    make_book(tmp_path / "farm-copy.xlsx", [record()])
    catalog, sync, settings = catalog_sync(tmp_path)
    sync.synchronize(settings)
    farms = catalog.search("Fazenda")
    assert len(farms) == 2
    assert len({farm.external_id for farm in farms}) == 2
    assert len({farm.parcels[0].external_id for farm in farms}) == 2
    catalog.close()


def test_catalog_search_covers_more_than_one_hundred_farms(tmp_path) -> None:
    make_book(tmp_path / "large.xlsx", [
        record(name=f"Fazenda {index:03d}", registration=f"MAT-{index:03d}")
        for index in range(105)
    ])
    catalog, sync, settings = catalog_sync(tmp_path)
    sync.synchronize(settings)
    assert len(catalog.search("Fazenda")) == 105
    catalog.close()


@pytest.mark.parametrize("term", [
    "Fazenda Exemplo", "MAT-1", "Pessoa Exemplo", "ANT-1", "Lote A",
    "Palmas", "CCIR-1", "ITR-1", "CAR-1", "000.000.000-00",
])
def test_catalog_searches_across_property_and_parcel_fields(tmp_path, term) -> None:
    make_book(tmp_path / "search.xlsx", [record()])
    catalog, sync, settings = catalog_sync(tmp_path)
    sync.synchronize(settings)
    assert len(catalog.search(term)) == 1
    catalog.close()


def test_document_mask_hides_middle_digits() -> None:
    masked = mask_document("000.000.000-00")
    assert masked != "000.000.000-00"
    assert "•" in masked
