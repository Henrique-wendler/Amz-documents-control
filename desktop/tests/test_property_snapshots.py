from __future__ import annotations

import sqlite3
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError

from amazon_agro.config.settings import AppSettings
from amazon_agro.domain.models import (
    Proposal, ProposalProperty, ProposalPropertyParcel, PropertyClassification,
)
from amazon_agro.exporters.excel import OpenpyxlExcelProposalExporter
from amazon_agro.repositories.sqlite_property_catalog import SQLitePropertyCatalogRepository
from amazon_agro.repositories.sqlite_proposals import SQLiteProposalRepository
from amazon_agro.services.export_validator import ProposalExportValidator
from amazon_agro.services.property_sync_service import PropertyCatalogSyncService
from amazon_agro.services.proposal_service import ProposalService
from test_property_workbooks import make_book, record


def _proposal() -> Proposal:
    return Proposal(
        numero_proposta="001", proponente="Produtor Exemplo",
        cpf_cnpj="000.000.000-00", tecnico="Técnica Exemplo",
        finalidade="Custeio", cidade="Palmas", valor_total=Decimal("1000"),
    )


def _link(proposal: Proposal, farm, selected=None) -> ProposalProperty:
    parcels = selected if selected is not None else farm.parcels
    return ProposalProperty(
        proposal_id=proposal.id, property_external_id=farm.external_id,
        classificacao=PropertyClassification.CLASS_1,
        property_name_snapshot=farm.name,
        municipality_snapshot=farm.municipality,
        source_file_snapshot=farm.source_file,
        owner_name_snapshot=farm.owner_name,
        selected_parcels=[ProposalPropertyParcel(
            parcel_external_id=parcel.external_id,
            registration_snapshot=parcel.registration,
            previous_registration_snapshot=parcel.previous_registration,
            area_snapshot=parcel.area,
            lot_description_snapshot=parcel.lot_description,
        ) for parcel in parcels],
    )


def test_multiple_selected_registrations_round_trip_in_proposal_snapshot(tmp_path) -> None:
    path = make_book(tmp_path / "farm.xlsx", [
        record(registration="MAT-1"),
        record(name=None, registration="MAT-2", **{
            "Proprietário": None, "CPF/CNPJ": None,
        }),
    ], merges=(("Fazenda", 1, 2), ("Proprietário", 1, 2), ("CPF/CNPJ", 1, 2)))
    catalog = SQLitePropertyCatalogRepository(tmp_path / "catalog.sqlite3")
    settings = AppSettings(property_source_directory=str(tmp_path))
    PropertyCatalogSyncService(catalog).synchronize(settings)
    farm = catalog.search("")[0]
    proposal = _proposal()
    proposal.properties.append(_link(proposal, farm))
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    ProposalService(repository, catalog).save(proposal)
    reopened = repository.get(proposal.id)
    assert reopened is not None
    assert [parcel.registration_snapshot for parcel in reopened.properties[0].selected_parcels] == [
        "MAT-1", "MAT-2"
    ]
    assert reopened.properties[0].property_name_snapshot == "Fazenda Exemplo"
    with repository.engine.connect() as connection:
        assert connection.exec_driver_sql("PRAGMA foreign_key_check").fetchall() == []
    assert ProposalExportValidator(settings, catalog).validate(reopened).ok
    output = OpenpyxlExcelProposalExporter(settings, catalog).export(
        reopened, tmp_path / "selected.xlsx"
    )
    sheet = load_workbook(output).active
    assert [sheet[f"I{row}"].value for row in (28, 29)] == ["MAT-1", "MAT-2"]
    catalog.close()
    repository.engine.dispose()


def test_historical_single_registration_exports_after_source_removal(tmp_path) -> None:
    source = make_book(tmp_path / "farm.xlsx", [record(registration="ORIGINAL")])
    catalog = SQLitePropertyCatalogRepository(tmp_path / "catalog.sqlite3")
    settings = AppSettings(property_source_directory=str(tmp_path))
    sync = PropertyCatalogSyncService(catalog)
    sync.synchronize(settings)
    farm = catalog.search("")[0]
    proposal = _proposal()
    proposal.properties.append(_link(proposal, farm))
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    ProposalService(repository, catalog).save(proposal)
    source.unlink()
    sync.synchronize(settings)
    assert catalog.search("") == []
    catalog.close()
    disconnected = SQLitePropertyCatalogRepository(tmp_path / "empty.sqlite3")
    historical = repository.get(proposal.id)
    assert historical is not None
    assert ProposalExportValidator(settings, disconnected).validate(historical).ok
    output = tmp_path / "historical.xlsx"
    OpenpyxlExcelProposalExporter(settings, disconnected).export(historical, output)
    sheet = load_workbook(output).active
    assert (sheet["A28"].value, sheet["E28"].value, sheet["I28"].value) == (
        "Fazenda Exemplo", "Palmas", "ORIGINAL"
    )
    disconnected.close()
    repository.engine.dispose()


def test_proposal_snapshot_survives_changed_registration(tmp_path) -> None:
    source = make_book(tmp_path / "farm.xlsx", [record(registration="BEFORE")])
    catalog = SQLitePropertyCatalogRepository(tmp_path / "catalog.sqlite3")
    settings = AppSettings(property_source_directory=str(tmp_path))
    sync = PropertyCatalogSyncService(catalog)
    sync.synchronize(settings)
    proposal = _proposal()
    proposal.properties.append(_link(proposal, catalog.search("")[0]))
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    ProposalService(repository, catalog).save(proposal)
    make_book(source, [record(registration="AFTER")])
    sync.synchronize(settings)
    historical = repository.get(proposal.id)
    assert historical.properties[0].selected_parcels[0].registration_snapshot == "BEFORE"
    output = tmp_path / "unchanged-proposal.xlsx"
    OpenpyxlExcelProposalExporter(settings, catalog).export(historical, output)
    assert load_workbook(output).active["I28"].value == "BEFORE"
    catalog.close()
    repository.engine.dispose()


def test_existing_proposal_database_receives_snapshot_columns(tmp_path) -> None:
    path = tmp_path / "legacy.sqlite3"
    connection = sqlite3.connect(path)
    connection.execute("""
        CREATE TABLE proposal_properties (
            proposal_id TEXT NOT NULL, property_external_id TEXT NOT NULL,
            classificacao TEXT NOT NULL,
            PRIMARY KEY (proposal_id, property_external_id)
        )
    """)
    connection.commit()
    connection.close()
    repository = SQLiteProposalRepository(path)
    with repository.engine.connect() as connection:
        columns = {row[1] for row in connection.exec_driver_sql(
            "PRAGMA table_info(proposal_properties)"
        )}
    assert {"property_name_snapshot", "municipality_snapshot", "state_snapshot",
            "source_file_snapshot", "owner_name_snapshot"} <= columns
    proposal = _proposal()
    repository.save(proposal)
    assert repository.get(proposal.id) is not None
    repository.engine.dispose()


def test_without_source_catalog_starts_empty(tmp_path) -> None:
    catalog = SQLitePropertyCatalogRepository(tmp_path / "empty.sqlite3")
    assert catalog.search("") == []
    assert catalog.stats().properties == 0
    catalog.close()


def test_selected_parcel_foreign_key_rejects_orphan(tmp_path) -> None:
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    with pytest.raises(IntegrityError):
        with repository.engine.begin() as connection:
            connection.execute(text("""
                INSERT INTO proposal_property_parcels (
                    proposal_id, property_external_id, parcel_external_id,
                    registration_snapshot, previous_registration_snapshot,
                    lot_description_snapshot
                ) VALUES ('missing', 'missing', 'parcel', 'MAT-1', '', '')
            """))
    repository.engine.dispose()


def test_source_settings_persist_directory_patterns_and_profile(tmp_path) -> None:
    config = tmp_path / "config.json"
    settings = AppSettings(
        property_source_directory="planilhas", property_file_patterns=["*.xlsm"],
        property_recursive=True, property_profile_name="Família Exemplo",
        property_profile_options={"header_scan_rows": 12},
    )
    settings.save(config)
    loaded = AppSettings.load(config)
    assert loaded.source_directory() == tmp_path / "planilhas"
    assert loaded.property_file_patterns == ["*.xlsm"]
    assert loaded.property_recursive is True
    assert loaded.property_profile_name == "Família Exemplo"
    assert loaded.property_profile_options == {"header_scan_rows": 12}
    assert "property_spreadsheet_path" not in config.read_text(encoding="utf-8")
    assert '"source_directory"' in config.read_text(encoding="utf-8")
    assert '"property_source_directory"' not in config.read_text(encoding="utf-8")


def test_legacy_source_directory_key_loads_and_saves_with_new_name(tmp_path) -> None:
    config = tmp_path / "config.json"
    settings = AppSettings()
    settings.save(config)
    raw = config.read_text(encoding="utf-8").replace(
        '"source_directory": ""', '"property_source_directory": "planilhas"'
    )
    config.write_text(raw, encoding="utf-8")
    loaded = AppSettings.load(config)
    assert loaded.source_directory() == tmp_path / "planilhas"
    loaded.save()
    assert '"source_directory": "planilhas"' in config.read_text(encoding="utf-8")
