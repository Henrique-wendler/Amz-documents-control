"""Synthetic layout regressions; no customer names, IDs or workbooks."""

from decimal import Decimal
from hashlib import sha256

import pytest
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from amazon_agro.integrations.workbook_inspector import NeedsConfigurationError
from amazon_agro.integrations.workbook_profile import PropertyWorkbookProfile
from test_property_workbooks import catalog_sync, make_book, parse, record


@pytest.mark.parametrize("area", ["178,9911ha", "178,9911", 178.9911, 1e-7, None])
def test_optional_area_forms_in_workbook(tmp_path, area):
    path = make_book(tmp_path / "areas.xlsx", [record(**{"Área (ha)": area})])
    parsed = parse(path)
    assert parsed.properties[0].parcels[0].area == (None if area is None else
                                                  Decimal("0.0000001") if area == 1e-7 else
                                                  Decimal("178.9911"))
    assert parsed.warnings == ()


def test_only_structural_columns_are_required(tmp_path):
    path = make_book(tmp_path / "minimal.xlsx", [record()], headers=("Fazenda", "Matrículas"))
    farm = parse(path).properties[0]
    assert farm.owner_name == farm.owner_document == farm.municipality == farm.state == ""
    assert farm.parcels[0].area is None


def test_configured_optional_requirement_is_still_enforced(tmp_path):
    path = make_book(tmp_path / "strict.xlsx", [record()], headers=("Fazenda", "Matrículas"))
    profile = PropertyWorkbookProfile.configured("Família", {
        "required_fields": ["name", "registration", "owner_name"]
    })
    with pytest.raises(NeedsConfigurationError, match="Cabeçalhos"):
        parse(path, profile)


def test_merged_registration_creates_one_parcel_with_later_complement(tmp_path):
    path = make_book(tmp_path / "complement.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Matrículas": "MAT-1"},
        {"Área (ha)": "178,9911ha", "Lote/Gleba": "Lote A", "CAR": "CAR-A"},
        {},
    ], merges=(("Fazenda", 1, 3), ("Matrículas", 1, 3)))
    farm = parse(path).properties[0]
    assert len(farm.parcels) == 1
    assert farm.parcels[0].area == Decimal("178.9911")
    assert farm.parcels[0].lot_description == "Lote A"
    assert farm.parcels[0].extra_fields["car"] == "CAR-A"


def test_farm_identification_precedes_registration_in_real_merge(tmp_path):
    path = make_book(tmp_path / "delayed.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Proprietário": "Pessoa Exemplo", "Município": "Cidade Exemplo"},
        {}, {"Matrículas": "MAT-1"}, {}, {"Matrículas": "MAT-2"},
    ], merges=(("Fazenda", 1, 5),))
    parsed = parse(path)
    assert len(parsed.properties) == 1
    farm = parsed.properties[0]
    assert len(farm.parcels) == 2
    assert farm.owner_name == "Pessoa Exemplo"
    assert farm.municipality == "Cidade Exemplo"
    assert not parsed.warnings


def test_unmerged_farm_before_registration_does_not_authorize_forward_fill(tmp_path):
    path = make_book(tmp_path / "unmerged.xlsx", [
        {"Fazenda": "Fazenda Alfa"}, {"Matrículas": "MAT-1"},
    ])
    with pytest.raises(NeedsConfigurationError, match="sem fazenda"):
        parse(path)


@pytest.mark.parametrize("field", ["Área (ha)", "CCIR", "ITR", "CAR", "Lote/Gleba", "Matrícula Anterior"])
def test_parcel_data_without_registration_cannot_be_silently_attached(tmp_path, field):
    path = make_book(tmp_path / "orphan.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Matrículas": "MAT-1"},
        {field: "12"}, {"Matrículas": "MAT-2"},
    ], merges=(("Fazenda", 1, 3),))
    with pytest.raises(NeedsConfigurationError, match="Linha 5 possui"):
        parse(path)


def test_shared_metadata_anchor_before_registration_uses_explicit_merge(tmp_path):
    path = make_book(tmp_path / "shared.xlsx", [
        {"Fazenda": "Fazenda Alfa", "CCIR": "CCIR-A"},
        {"Matrículas": "MAT-1"}, {"Matrículas": "MAT-2"},
    ], merges=(("Fazenda", 1, 3), ("CCIR", 1, 3)))
    farm = parse(path).properties[0]
    assert farm.ccir == "CCIR-A"
    assert [p.extra_fields["ccir"] for p in farm.parcels] == ["CCIR-A", "CCIR-A"]


def test_area_merge_covering_one_registration_is_read_once(tmp_path):
    path = make_book(tmp_path / "single-area.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Área (ha)": "178,9911ha"},
        {"Matrículas": "MAT-1"}, {},
    ], merges=(("Fazenda", 1, 3), ("Área (ha)", 1, 3)))
    assert parse(path).properties[0].total_area == Decimal("178.9911")


def test_registration_merge_cannot_cross_farms(tmp_path):
    path = make_book(tmp_path / "crossing.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Matrículas": "MAT-1"},
        {"Fazenda": "Fazenda Beta"},
    ], merges=(("Matrículas", 1, 2),))
    with pytest.raises(NeedsConfigurationError, match="atravessa grupos"):
        parse(path)


def test_same_name_in_separate_physical_merges_is_not_collapsed(tmp_path):
    path = make_book(tmp_path / "separate.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Matrículas": "MAT-1"}, {},
        {"Fazenda": "Fazenda Alfa", "Matrículas": "MAT-2"}, {},
    ], merges=(("Fazenda", 1, 2), ("Fazenda", 3, 4)))
    with pytest.raises(NeedsConfigurationError, match="Blocos separados"):
        parse(path)


def test_multiple_owners_in_registration_merge_are_not_discarded(tmp_path):
    path = make_book(tmp_path / "owners.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Matrículas": "MAT-1", "Proprietário": "Pessoa Alfa"},
        {"Proprietário": "Pessoa Beta"}, {"Proprietário": "Pessoa Gama"},
        {"Proprietário": "Pessoa Delta"},
    ], merges=(("Fazenda", 1, 4), ("Matrículas", 1, 4)))
    farm = parse(path).properties[0]
    assert len(farm.parcels) == 1
    assert farm.owner_count == len(farm.parcels[0].owner_links) == 4
    assert farm.owner_name == farm.owner_document == ""


def test_different_owners_do_not_split_one_physical_farm(tmp_path):
    path = make_book(tmp_path / "farm-owners.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Matrículas": "MAT-1", "Proprietário": "Pessoa Alfa"},
        {"Matrículas": "MAT-2", "Proprietário": "Pessoa Beta"},
    ], merges=(("Fazenda", 1, 2),))
    farm = parse(path).properties[0]
    assert len(farm.parcels) == farm.owner_count == 2
    assert set(farm.parcels[0].owner_ids).isdisjoint(farm.parcels[1].owner_ids)


def test_registration_complement_conflict_is_diagnosed(tmp_path):
    path = make_book(tmp_path / "conflict.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Matrículas": "MAT-1", "CAR": "CAR-A"},
        {"CAR": "CAR-B"},
    ], merges=(("Fazenda", 1, 2), ("Matrículas", 1, 2)))
    with pytest.raises(NeedsConfigurationError, match="CAR conflitante"):
        parse(path)


def test_different_parcel_documents_survive_sync_and_search(tmp_path):
    path = make_book(tmp_path / "documents.xlsx", [
        record(**{"CCIR": "CCIR-A", "ITR": "ITR-A", "CAR": "CAR-A"}),
        record(name=None, registration="MAT-2", **{"CCIR": "CCIR-B", "ITR": "ITR-B", "CAR": "CAR-B"}),
    ], merges=(("Fazenda", 1, 2),))
    before = sha256(path.read_bytes()).digest()
    catalog, sync, settings = catalog_sync(tmp_path)
    try:
        assert sync.synchronize(settings).events[0].status == "OK"
        farm = catalog.search("CCIR-B")[0]
        assert farm.ccir == farm.itr == farm.car == ""
        assert farm.parcels[0].extra_fields == {"ccir": "CCIR-A", "itr": "ITR-A", "car": "CAR-A"}
        assert farm.parcels[1].extra_fields == {"ccir": "CCIR-B", "itr": "ITR-B", "car": "CAR-B"}
        assert all(p.property_external_id == farm.external_id for p in farm.parcels)
        assert catalog.connection.execute("PRAGMA foreign_key_check").fetchall() == []
        assert sha256(path.read_bytes()).digest() == before
    finally:
        catalog.close()


def test_notes_and_empty_rows_do_not_create_entities_and_warnings_persist(tmp_path):
    path = make_book(tmp_path / "notes.xlsx", [record(), {}, {"Fazenda": "Observação: revisar documentos"}])
    workbook = load_workbook(path)
    workbook.active["N5"] = "Observação externa"
    workbook.save(path)
    parsed = parse(path)
    assert len(parsed.properties) == len(parsed.properties[0].parcels) == 1
    assert len(parsed.warnings) == 2
    catalog, sync, settings = catalog_sync(tmp_path)
    try:
        report = sync.synchronize(settings)
        assert report.events[0].status == "OK" and report.stats.warnings == 1
        warning = catalog.get_source(path).last_error
        assert "N5" in warning and "Linha 6" in warning
        assert sync.synchronize(settings).events[0].action == "UNCHANGED"
        assert catalog.get_source(path).last_error == warning
    finally:
        catalog.close()


@pytest.mark.parametrize("registration", ["Não há", "-", "0", "ver observação"])
def test_invalid_registration_never_creates_empty_or_fake_parcel(tmp_path, registration):
    path = make_book(tmp_path / "invalid-id.xlsx", [record(registration=registration)])
    with pytest.raises(NeedsConfigurationError, match="matrícula não identificável"):
        parse(path)


def test_invalid_area_diagnostic_never_discloses_document(tmp_path):
    path = make_book(tmp_path / "invalid-area.xlsx", [record(**{"Área (ha)": "não informado"})])
    with pytest.raises(NeedsConfigurationError, match="Linha 4: Área inválida") as error:
        parse(path)
    assert "000.000.000-00" not in str(error.value)


def test_formulas_are_not_mistaken_for_empty_optional_values(tmp_path):
    path = make_book(tmp_path / "formula.xlsx", [record(**{"Área (ha)": "=1+1"})])
    with pytest.raises(NeedsConfigurationError, match="fórmula ou erro Excel"):
        parse(path)


def test_fill_colors_have_no_semantic_effect(tmp_path):
    path = make_book(tmp_path / "colors.xlsx", [record()])
    before = parse(path)
    workbook = load_workbook(path)
    for cell in workbook.active[4]:
        cell.fill = PatternFill("solid", fgColor="FFFF00" if cell.column % 2 else "00FF00")
    workbook.save(path)
    assert parse(path) == before


def test_parser_revision_reprocesses_previously_valid_catalog(tmp_path):
    path = make_book(tmp_path / "version.xlsx", [record()])
    catalog, sync, settings = catalog_sync(tmp_path)
    try:
        sync.synchronize(settings)
        catalog.connection.execute("UPDATE source_files SET profile_signature=?", ("old-parser",))
        catalog.connection.commit()
        assert sync.synchronize(settings).events[0].action == "MODIFIED"
    finally:
        catalog.close()


def test_ambiguous_update_keeps_previous_catalog_without_partial_import(tmp_path):
    path = make_book(tmp_path / "atomic.xlsx", [record()])
    catalog, sync, settings = catalog_sync(tmp_path)
    try:
        sync.synchronize(settings)
        previous = catalog.search("")[0]
        make_book(path, [record(), {"CAR": "CAR-ORPHAN"}])
        report = sync.synchronize(settings)
        assert report.events[0].status == "NEEDS_CONFIGURATION"
        assert catalog.search("")[0].external_id == previous.external_id
        assert catalog.stats().properties == catalog.stats().parcels == 1
        assert "Linha 5" in catalog.get_source(path).last_error
    finally:
        catalog.close()


def test_missing_optional_document_is_not_a_second_owner(tmp_path):
    path = make_book(tmp_path / "partial-owner.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Matrículas": "MAT-1", "Proprietário": "Pessoa Alfa",
         "CPF/CNPJ": "000.000.000-00"}, {},
    ], merges=(("Fazenda", 1, 2), ("Matrículas", 1, 2), ("Proprietário", 1, 2)))
    farm = parse(path).properties[0]
    assert len(farm.parcels) == 1
    assert farm.owner_document == "000.000.000-00"


def test_missing_parcel_metadata_does_not_inherit_another_parcels_value(tmp_path):
    path = make_book(tmp_path / "metadata-gap.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Matrículas": "MAT-1", "CAR": "CAR-A"},
        {"Matrículas": "MAT-2"},
    ], merges=(("Fazenda", 1, 2),))
    farm = parse(path).properties[0]
    assert farm.car == ""
    assert farm.parcels[0].extra_fields["car"] == "CAR-A"
    assert "car" not in farm.parcels[1].extra_fields


def test_text_areas_and_unequal_merge_heights_across_nine_registrations(tmp_path):
    rows = []
    merges = []
    for index in range(9):
        first = len(rows) + 1
        height = 4 if index < 8 else 2
        rows.append({"Fazenda": f"Fazenda {index // 4}", "Matrículas": f"MAT-{index + 1}",
                     "Área (ha)": "178,9911ha", "Proprietário": "Pessoa Exemplo",
                     "CCIR": f"CCIR-{index // 3}", "CAR": f"CAR-{index // 4}"})
        rows.extend({} for _ in range(height - 1))
        merges.append(("Matrículas", first, first + height - 1))
        # An area merge may be shorter than its single registration merge.
        merges.append(("Área (ha)", first, first + height - 2 if height == 4 else first + 1))
        merges.append(("Proprietário", first, first + height - 1))
    merges.extend((("Fazenda", 1, 16), ("Fazenda", 17, 32), ("Fazenda", 33, 34)))
    path = make_book(tmp_path / "text-layout.xlsx", rows, merges=tuple(merges))
    parsed = parse(path)
    assert len(parsed.properties) == 3
    assert sum(len(p.parcels) for p in parsed.properties) == 9
    assert sum(p.total_area for p in parsed.properties) == Decimal("178.9911") * 9
    assert not parsed.warnings


def test_full_multi_owner_layout_reports_all_28_complement_rows(tmp_path):
    headers = ("Fazenda", "Matrículas", "Área (ha)", "Proprietário", "CPF/CNPJ", "CCIR", "ITR", "CAR")
    rows = []
    merges = []
    for index in range(9):
        first = len(rows) + 1
        for owner in range(4):
            rows.append({
                "Fazenda": f"Fazenda {index // 7}" if owner == 0 else None,
                "Matrículas": f"MAT-{index + 1}" if owner == 0 else None,
                "Área (ha)": 10.25 if owner == 0 else None,
                "Proprietário": f"Pessoa {owner}",
                "CPF/CNPJ": f"00000000{owner:03d}",
                "CCIR": f"CCIR-{index}" if owner == 0 else None,
            })
        merges.extend((key, first, first + 3) for key in ("Matrículas", "Área (ha)", "CCIR"))
    for index in range(3):
        rows.append({"Fazenda": "Fazenda 2", "Matrículas": f"MAT-{index + 10}",
                     "Área (ha)": 10.25, "Proprietário": f"Pessoa {index + 4}",
                     "CPF/CNPJ": f"00000000{index + 4:03d}"})
    rows.extend((
        {"Fazenda": "Fazenda 3", "Matrículas": "MAT-13", "Área (ha)": 10.25,
         "Proprietário": "Pessoa 7", "CPF/CNPJ": "00000000007"},
        {"Proprietário": "Pessoa 8", "CPF/CNPJ": "00000000008"},
    ))
    merges.extend((("Fazenda", 1, 28), ("Fazenda", 29, 36), ("Fazenda", 37, 39),
                   ("Fazenda", 40, 41), ("Matrículas", 40, 41), ("Área (ha)", 40, 41)))
    path = make_book(tmp_path / "multi-owner-layout.xlsx", rows, header_row=7,
                     headers=headers, merges=tuple(merges))
    workbook = load_workbook(path)
    for row in (10, 31, 33, 34):
        workbook.active.cell(row, 10, "Anotação externa")
    workbook.save(path)
    before = sha256(path.read_bytes()).digest()
    parsed = parse(path)
    assert len(parsed.properties) == 4
    assert sum(len(f.parcels) for f in parsed.properties) == 13
    assert len(parsed.owners) == 9
    assert sum(len(p.owner_links) for f in parsed.properties for p in f.parcels) == 41
    assert [f.owner_count for f in parsed.properties] == [4, 4, 3, 2]
    assert len(parsed.warnings) == 4
    assert sha256(path.read_bytes()).digest() == before
