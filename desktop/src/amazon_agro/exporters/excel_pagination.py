"""Template-preserving continuation blocks and explicit print boundaries."""
from copy import copy

from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.worksheet import Worksheet

from amazon_agro.exporters.excel_map import PDF_PRINT_AREA, PROPERTY_ROWS
from amazon_agro.exporters.excel_layout import add_logo


def _copy_rows(sheet: Worksheet, first: int, last: int, destination: int) -> None:
    offset = destination - first
    merges = list(sheet.merged_cells.ranges)
    for row in range(first, last + 1):
        target_row = row + offset
        dimension = copy(sheet.row_dimensions[row])
        dimension.index = target_row
        sheet.row_dimensions[target_row] = dimension
        for column in range(1, 14):
            source = sheet.cell(row, column)
            target = sheet.cell(target_row, column)
            target.value = source.value
            target._style = copy(source._style)
            if source.comment is not None:
                target.comment = copy(source.comment)
    for merged in merges:
        if first <= merged.min_row and merged.max_row <= last:
            sheet.merge_cells(
                start_row=merged.min_row + offset, end_row=merged.max_row + offset,
                start_column=merged.min_col, end_column=merged.max_col,
            )


def property_page_rows(sheet: Worksheet, line_count: int) -> list[tuple[int, ...]]:
    """Return four data rows per page; never fit all pages to one page tall."""
    page_count = max(1, (line_count + len(PROPERTY_ROWS) - 1) // len(PROPERTY_ROWS))
    pages = [PROPERTY_ROWS]
    areas = [PDF_PRINT_AREA]
    previous_end = 35
    for _ in range(1, page_count):
        start = previous_end + 1
        # Identity and proponent repeat; only the property section continues.
        _copy_rows(sheet, 1, 6, start)
        add_logo(sheet, start)
        _copy_rows(sheet, 26, 35, start + 6)
        sheet.cell(start + 6, 1, "IV - IMÓVEIS (CONTINUAÇÃO)")
        pages.append(tuple(range(start + 8, start + 12)))
        end = start + 15
        areas.append(f"A{start}:M{end}")
        sheet.row_breaks.append(Break(id=previous_end))
        previous_end = end
    sheet.print_area = areas
    if page_count > 1:
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.fitToWidth = 1
        sheet.oddFooter.right.text = "Página &P de &N"
    return pages
