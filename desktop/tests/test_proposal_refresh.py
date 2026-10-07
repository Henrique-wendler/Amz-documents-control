"""Approved FNO-only ABC rule and controlled desktop refresh, synthetic data only."""
from decimal import Decimal
from pathlib import Path
import os
import sqlite3
import subprocess
import sys

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QPoint, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QLabel, QRadioButton

from amazon_agro.config.settings import AppSettings
from amazon_agro.domain.models import Proposal
from amazon_agro.exporters.excel import OpenpyxlExcelProposalExporter
from amazon_agro.exporters.pdf_backends import ExcelComPdfBackend
from amazon_agro.repositories.fake_properties import FakePropertyRepository
from amazon_agro.repositories.sqlite_proposals import SQLiteProposalRepository
from amazon_agro.services.export_service import ExportValidationError
from amazon_agro.ui.boolean_choice import BooleanChoice
from amazon_agro.ui.currency_input import CurrencyInput
from amazon_agro.ui.pages import ProposalPage
from amazon_agro.ui.review import ReviewPage
from amazon_agro.services.proposal_service import ProposalService
from test_business_definitions import proposal_with_parcels
from test_hotfix_qa import app
from test_proposal_ui import fill, ui


@pytest.mark.parametrize('fno,percent,expected', [
    ('70000', '5', '3500'), ('0', '5', '0'), ('100000', '0', '0'),
    ('100000.01', '3.25', '3250.000325'), ('123.45', '2.99', '3.691155'),
])
def test_abc_calculates_from_fno_with_decimal_without_intermediate_rounding(fno, percent, expected):
    proposal = Proposal(valor_total=Decimal('200000'), valor_fno=Decimal(fno), valor_of=Decimal('30000'),
                        laudo_abc_financiado=True, laudo_abc_percentual=Decimal(percent),
                        laudo_abc_valor=Decimal('999999'))
    assert proposal.calculated_abc_amount == Decimal(expected)
    assert isinstance(proposal.calculated_abc_amount, Decimal)


def test_abc_no_has_no_effective_amount():
    proposal = Proposal(valor_total=Decimal('100000'), laudo_abc_percentual=Decimal('5'),
                        laudo_abc_valor=Decimal('9999'))
    assert proposal.calculated_abc_amount is None


def test_abc_ui_changes_fno_and_percentage_immediately_and_total_is_not_its_basis(app):
    page = ProposalPage()
    page.abc_yes.setChecked(True)
    page.spins['valor_total'].setValue(Decimal('100000'))
    page.spins['valor_fno'].setValue(Decimal('70000'))
    page.spins['valor_of'].setValue(Decimal('30000'))
    page.spins['laudo_abc_percentual'].setValue(5)
    assert page.abc_amount.text() == 'R$ 3.500,00'
    page.spins['valor_total'].setValue(Decimal('120000'))
    assert page.abc_amount.text() == 'R$ 3.500,00'
    page.spins['valor_fno'].setValue(Decimal('80000'))
    assert page.abc_amount.text() == 'R$ 4.000,00'
    page.spins['laudo_abc_percentual'].setValue(2.5)
    assert page.abc_amount.text() == 'R$ 2.000,00'
    assert isinstance(page.abc_amount, QLabel)
    assert 'laudo_abc_valor' not in page.spins
    page.abc_no.setChecked(True)
    assert not page.abc_yes.isChecked()
    assert page.abc_fields.isHidden() and not page.abc_fields.isEnabled()
    page.close()


def test_abc_calculated_value_persists_and_reopening_recalculates_legacy_amount(app, tmp_path):
    page = ProposalPage()
    proposal = Proposal(valor_total=Decimal('100000'), valor_fno=Decimal('70000'), laudo_abc_financiado=True,
                        laudo_abc_percentual=Decimal('5'), laudo_abc_valor=Decimal('1'))
    page.load(proposal)
    assert page.abc_amount.text() == 'R$ 3.500,00'
    page.read_into(proposal)
    repository = SQLiteProposalRepository(tmp_path / 'draft.sqlite3')
    repository.save(proposal)
    repository.engine.dispose()
    repository = SQLiteProposalRepository(tmp_path / 'draft.sqlite3')
    reopened = repository.get(proposal.id)
    assert reopened.laudo_abc_valor == Decimal('3500')
    assert reopened.calculated_abc_amount == Decimal('3500')
    page.load(reopened)
    assert page.abc_amount.text() == 'R$ 3.500,00'
    page.close()
    repository.engine.dispose()


def test_own_resources_ui_has_only_yes_no_and_conditional_percentage(app):
    page = ProposalPage()
    assert 'recursos_proprios' not in page.spins
    assert len(page.findChildren(CurrencyInput)) == 3
    assert page.own_resources_fields.isHidden()
    page.own_resources_yes.setChecked(True)
    assert not page.own_resources_no.isChecked()
    assert not page.own_resources_fields.isHidden()
    page.spins['percentual_recursos_proprios'].setValue(10)
    proposal = Proposal(recursos_proprios=Decimal('123.45'))
    page.read_into(proposal)
    assert proposal.has_own_resources and proposal.own_resources_percentage == Decimal('10')
    page.own_resources_no.setChecked(True)
    assert page.own_resources_fields.isHidden() and not page.own_resources_fields.isEnabled()
    page.read_into(proposal)
    assert not proposal.has_own_resources and proposal.own_resources_percentage == 0
    assert proposal.percentual_recursos_proprios == 0
    assert proposal.recursos_proprios == Decimal('123.45')  # Legacy monetary data retained.
    page.close()


@pytest.mark.parametrize('enabled,percent', [(True, '0'), (True, '10'), (False, '0')])
def test_own_resources_explicit_choice_survives_save_reopen_even_when_zero(app, tmp_path, enabled, percent):
    page = ProposalPage()
    proposal = Proposal(recursos_proprios=Decimal('500'))
    page.load(proposal)
    page.own_resources_choice.setValue(enabled)
    page.spins['percentual_recursos_proprios'].setValue(float(percent))
    page.read_into(proposal)
    repository = SQLiteProposalRepository(tmp_path / 'own.sqlite3')
    repository.save(proposal)
    repository.engine.dispose()
    repository = SQLiteProposalRepository(tmp_path / 'own.sqlite3')
    reopened = repository.get(proposal.id)
    assert reopened.has_own_resources is enabled
    assert reopened.possui_recursos_proprios is enabled
    assert reopened.own_resources_percentage == Decimal(percent)
    page.load(reopened)
    assert page.own_resources_yes.isChecked() is enabled
    assert page.own_resources_fields.isHidden() is not enabled
    page.close()
    repository.engine.dispose()


def test_own_resources_nullable_migration_is_idempotent_and_preserves_legacy_snapshots(tmp_path):
    path = tmp_path / 'legacy.sqlite3'
    repository = SQLiteProposalRepository(path)
    proposal = proposal_with_parcels(5)
    repository.save(proposal)
    baseline_properties = repository.get(proposal.id).properties
    repository.engine.dispose()
    with sqlite3.connect(path) as connection:
        connection.execute('ALTER TABLE proposals DROP COLUMN possui_recursos_proprios')
    for _ in range(2):
        repository = SQLiteProposalRepository(path)
        reopened = repository.get(proposal.id)
        assert reopened.possui_recursos_proprios is None and reopened.has_own_resources
        assert reopened.recursos_proprios == Decimal('100')
        assert reopened.properties == baseline_properties
        repository.engine.dispose()
    with sqlite3.connect(path) as connection:
        assert connection.execute('PRAGMA integrity_check').fetchone() == ('ok',)
        assert connection.execute('PRAGMA foreign_key_check').fetchall() == []


@pytest.mark.parametrize('field', ['astec_fno_financiada', 'astec_of_financiada'])
def test_astec_uses_exclusive_keyboard_accessible_yes_no_and_never_hides_financing(app, field):
    page = ProposalPage()
    page.show()
    choice = page.flags[field]
    assert isinstance(choice, BooleanChoice)
    assert isinstance(choice.yes, QRadioButton) and choice.no.isChecked()
    choice.yes.setFocus()
    QTest.keyClick(choice.yes, Qt.Key.Key_Space)
    assert choice.value() and not choice.no.isChecked()
    percent_name = field.replace('financiada', 'percentual')
    assert page.spins[percent_name].isVisible() and page.spins[percent_name].isEnabled()
    choice.no.setChecked(True)
    for name in ('valor_fno', 'valor_of'):
        assert page.spins[name].isVisible() and page.spins[name].isEnabled()
    assert page.astec_fields[field].isHidden() and not page.spins[percent_name].isEnabled()
    page.spins['valor_total'].setValue(Decimal('100000'))
    page.spins['valor_fno'].setValue(Decimal('70000'))
    page.spins['valor_of'].setValue(Decimal('30000'))
    assert page.financing_shares['valor_fno'].text() == 'Participação FNO: 70,00%'
    assert page.financing_shares['valor_of'].text() == 'Participação OF: 30,00%'
    proposal = Proposal()
    page.read_into(proposal)
    assert getattr(proposal, field) is False
    page.close()


@pytest.mark.parametrize('fno,of', [('100001', '0'), ('0', '100001'), ('70000', '40000')])
def test_financial_inconsistency_saves_but_status_and_export_are_blocked(ui, tmp_path, fno, of):
    window, app = ui
    fill(window)
    page = window.proposal_page
    page.spins['valor_total'].setValue(Decimal('100000'))
    page.spins['valor_fno'].setValue(Decimal(fno))
    page.spins['valor_of'].setValue(Decimal(of))
    app.processEvents()
    validation = window.export_validator.validate(window._collect())
    assert not validation.ok and validation.financial_errors
    assert window.badge.property('ready') is False and window.badge.property('blocked') is True
    assert 'Pronta para gerar' not in window.badge.text()
    assert window.summary.status.property('blocked') is True
    window.save_proposal()
    assert window.service.get(window.current.id) is not None
    with pytest.raises(ExportValidationError):
        window.export_service.generate_excel(window.current.id, tmp_path / 'exports')
    assert not list((tmp_path / 'exports').glob('*.xlsx'))


@pytest.mark.parametrize('fno,of', [('70000', '30000'), ('60000', '25000')])
def test_valid_financial_composition_is_ready_and_exports(ui, tmp_path, fno, of):
    window, app = ui
    fill(window)
    page = window.proposal_page
    page.spins['valor_total'].setValue(Decimal('100000'))
    page.spins['valor_fno'].setValue(Decimal(fno))
    page.spins['valor_of'].setValue(Decimal(of))
    page.own_resources_yes.setChecked(True)
    page.spins['percentual_recursos_proprios'].setValue(10)
    app.processEvents()
    assert window.badge.property('ready') is True and window.badge.property('blocked') is False
    window.save_proposal()
    (tmp_path / 'exports').mkdir()
    result = window.export_service.generate_excel(window.current.id, tmp_path / 'exports')
    assert result.xlsx_path.is_file()


@pytest.mark.parametrize('enabled', [False, True])
def test_xlsx_keeps_layout_and_exports_calculated_abc_and_percentage_only_resources(tmp_path, enabled):
    proposal = proposal_with_parcels(1)
    proposal.valor_total = Decimal('100000')
    proposal.valor_fno = Decimal('70000')
    proposal.valor_of = Decimal('30000')
    proposal.laudo_abc_financiado = enabled
    proposal.laudo_abc_percentual = Decimal('5')
    proposal.laudo_abc_valor = Decimal('9999')
    proposal.possui_recursos_proprios = enabled
    proposal.percentual_recursos_proprios = Decimal('10')
    path = OpenpyxlExcelProposalExporter(AppSettings(), FakePropertyRepository()).export(proposal, tmp_path / 'proposal.xlsx')
    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet['I23'].value == ('Sim' if enabled else 'Não')
    assert sheet['K23'].value == (0.05 if enabled else None)
    assert sheet['I25'].value == (3500 if enabled else None)
    assert sheet['G21'].value == ('Sim' if enabled else 'Não')
    assert sheet['H21'].value == (0.1 if enabled else 0)
    assert len(sheet._images) == 1 and sheet.page_setup.fitToHeight == 1
    workbook.close()


@pytest.mark.parametrize('size', [(1366, 768), (1920, 1080)])
def test_refresh_layout_contains_inputs_controls_and_bottom_actions_without_horizontal_scroll(ui, size):
    window, app = ui
    fill(window)
    window.resize(*size)
    page = window.proposal_page
    page.own_resources_yes.setChecked(True)
    page.abc_yes.setChecked(True)
    for choice in page.flags.values():
        choice.setValue(True)
    window.steps.setCurrentRow(2)
    app.processEvents()
    QTest.qWait(150)
    assert (window.width(), window.height()) == size
    scroll = window.stack.widget(2)
    assert not scroll.horizontalScrollBar().isVisible()
    assert page.minimumSizeHint().width() <= scroll.viewport().width()
    controls = list(page.spins.values()) + page.findChildren(QRadioButton) + [page.abc_amount]
    for control in controls:
        center = control.mapTo(page, control.rect().center())
        scroll.ensureVisible(center.x(), center.y(), 0, control.height() // 2 + 8)
        app.processEvents()
        position = control.mapTo(scroll.viewport(), QPoint(0, 0))
        assert 0 <= position.x() and position.x() + control.width() <= scroll.viewport().width()
        assert 0 <= position.y() and position.y() + control.height() <= scroll.viewport().height()
        if isinstance(control, QRadioButton):
            assert control.width() >= control.fontMetrics().horizontalAdvance(control.text()) + 24
    for button in [*window.header_buttons.values(), window.summary_toggle]:
        position = button.mapTo(window.centralWidget(), QPoint(0, 0))
        assert 0 <= position.x() and position.x() + button.width() <= window.centralWidget().width()
        assert 0 <= position.y() and position.y() + button.height() <= window.centralWidget().height()
    window.toggle_summary()
    assert window.summary_scroll.isHidden()
    window.toggle_summary()
    assert not window.summary_scroll.isHidden()


PDF_READER_PYTHON = os.environ.get('AMAZON_AGRO_PDF_READER_PYTHON')


@pytest.mark.skipif(not PDF_READER_PYTHON or not ExcelComPdfBackend().is_available(),
                    reason='Native Excel COM and pypdf reader required')
@pytest.mark.parametrize('enabled', [False, True])
def test_native_pdf_uses_fno_based_abc_and_conditional_astec_percentages(tmp_path, enabled):
    proposal = proposal_with_parcels(1)
    proposal.valor_total = Decimal('100000')
    proposal.valor_fno = Decimal('70000')
    proposal.valor_of = Decimal('30000')
    proposal.laudo_abc_financiado = enabled
    proposal.laudo_abc_percentual = Decimal('5')
    proposal.laudo_abc_valor = Decimal('9999')
    proposal.astec_fno_financiada = proposal.astec_of_financiada = enabled
    proposal.astec_fno_percentual = Decimal('2.75')
    proposal.astec_of_percentual = Decimal('1.25')
    xlsx = OpenpyxlExcelProposalExporter(AppSettings(), FakePropertyRepository()).export(proposal, tmp_path / 'abc.xlsx')
    pdf = tmp_path / 'abc.pdf'
    subprocess.run([sys.executable, '-c',
                    'from pathlib import Path; import sys; '
                    'from amazon_agro.exporters.pdf import SpreadsheetPdfProposalExporter; '
                    'from amazon_agro.exporters.pdf_backends import PdfBackendSelector, ExcelComPdfBackend; '
                    'SpreadsheetPdfProposalExporter(PdfBackendSelector([ExcelComPdfBackend()])).export(Path(sys.argv[1]), Path(sys.argv[2]))',
                    str(xlsx), str(pdf)], capture_output=True, text=True, check=True, timeout=120)
    read = subprocess.run([PDF_READER_PYTHON, '-c',
                           'from pypdf import PdfReader; import sys; r=PdfReader(sys.argv[1]); '
                           'assert len(r.pages)==1; print(r.pages[0].extract_text())', str(pdf)],
                          capture_output=True, text=True, check=True, timeout=30,
                          encoding="utf-8", env={**os.environ, "PYTHONIOENCODING": "utf-8"})
    text = ' '.join(read.stdout.split())
    assert 'LAUDO ABC FINANCIADO?' in text
    assert ('3.500,00' in text) is enabled
    assert ('2,75%' in text) is enabled
    assert ('1,25%' in text) is enabled
    assert 'ASTEC FNO FINANCIADA?' in text and 'ASTEC OF FINANCIADA?' in text
    assert 'SOMENTE FNO' in text
    assert '70,00%' in text and '30,00%' in text
    assert '9.999,00' not in text and '#' not in text


@pytest.mark.parametrize('fno_enabled,of_enabled', [(False, False), (True, False), (False, True), (True, True)])
def test_astec_choices_are_independent_persist_reopen_and_export_only_active_percentages(
    app, tmp_path, fno_enabled, of_enabled
):
    proposal = proposal_with_parcels(1)
    proposal.valor_total, proposal.valor_fno, proposal.valor_of = map(Decimal, ('100000', '70000', '30000'))
    proposal.astec_fno_percentual, proposal.astec_of_percentual = Decimal('2.75'), Decimal('1.25')
    proposal.laudo_abc_financiado, proposal.laudo_abc_percentual = True, Decimal('5')
    page = ProposalPage()
    page.load(proposal)
    for source, enabled in (('fno', fno_enabled), ('of', of_enabled)):
        flag = f'astec_{source}_financiada'
        page.flags[flag].setValue(enabled)
        assert page.astec_fields[flag].isHidden() is not enabled
        assert page.spins[f'astec_{source}_percentual'].isEnabled() is enabled
    page.read_into(proposal)
    repository = SQLiteProposalRepository(tmp_path / 'astec.sqlite3')
    repository.save(proposal)
    repository.engine.dispose()
    repository = SQLiteProposalRepository(tmp_path / 'astec.sqlite3')
    reopened = repository.get(proposal.id)
    assert reopened.astec_fno_financiada is fno_enabled
    assert reopened.astec_of_financiada is of_enabled
    assert reopened.astec_fno_percentual == Decimal('2.75')
    assert reopened.astec_of_percentual == Decimal('1.25')
    assert reopened.calculated_abc_amount == Decimal('3500')
    page.load(reopened)
    for source, enabled, percent in (('fno', fno_enabled, 2.75), ('of', of_enabled, 1.25)):
        assert page.flags[f'astec_{source}_financiada'].value() is enabled
        assert page.spins[f'astec_{source}_percentual'].value() == percent
        assert page.astec_fields[f'astec_{source}_financiada'].isHidden() is not enabled
    review = ReviewPage(ProposalService(repository, FakePropertyRepository()), AppSettings())
    review.load(reopened)
    assert ('% ASTEC FNO: 2,75%' in review.text.toPlainText()) is fno_enabled
    assert ('% ASTEC OF: 1,25%' in review.text.toPlainText()) is of_enabled
    xlsx = OpenpyxlExcelProposalExporter(AppSettings(), FakePropertyRepository()).export(reopened, tmp_path / 'astec.xlsx')
    workbook = load_workbook(xlsx)
    sheet = workbook.active
    assert sheet['F23'].value == ('Sim' if fno_enabled else 'Não')
    assert sheet['H23'].value == (0.0275 if fno_enabled else None)
    assert sheet['F25'].value == ('Sim' if of_enabled else 'Não')
    assert sheet['H25'].value == (0.0125 if of_enabled else None)
    assert sheet['I25'].value == 3500 and 'SOMENTE FNO' in sheet['I24'].value
    assert sheet['D23'].value == 0.7 and sheet['D25'].value == 0.3
    workbook.close()
    page.close()
    review.close()
    repository.engine.dispose()


@pytest.mark.parametrize('source', ['fno', 'of'])
@pytest.mark.parametrize('percent', [Decimal('-0.01'), Decimal('100.01'), Decimal('NaN'), Decimal('Infinity')])
def test_active_astec_percentage_requires_finite_decimal_between_zero_and_one_hundred(source, percent):
    proposal = Proposal(**{f'astec_{source}_financiada': True, f'astec_{source}_percentual': percent})
    with pytest.raises(ValueError, match=f'Percentual ASTEC {source.upper()}'):
        proposal.validate()


@pytest.mark.parametrize('source', ['fno', 'of'])
def test_inactive_legacy_astec_percentage_does_not_block_or_export(tmp_path, source):
    proposal = proposal_with_parcels(1)
    setattr(proposal, f'astec_{source}_percentual', Decimal('150'))
    proposal.validate()
    path = OpenpyxlExcelProposalExporter(AppSettings(), FakePropertyRepository()).export(proposal, tmp_path / 'inactive.xlsx')
    workbook = load_workbook(path)
    assert workbook.active['H23' if source == 'fno' else 'H25'].value is None
    workbook.close()


def test_abc_belongs_to_fno_even_in_partial_draft_with_total_zero():
    proposal = Proposal(valor_fno=Decimal('70000'), valor_total=Decimal('0'),
                        laudo_abc_financiado=True, laudo_abc_percentual=Decimal('5'))
    assert proposal.calculated_abc_amount == Decimal('3500')
    assert proposal.fno_percentage == 0 and proposal.of_percentage == 0


def test_proposal_groups_abc_and_astec_with_their_respective_financing_source(app):
    page = ProposalPage()
    def card(widget):
        parent = widget.parentWidget()
        while parent.objectName() != 'formCard':
            parent = parent.parentWidget()
        return parent
    fno = card(page.spins['valor_fno'])
    of = card(page.spins['valor_of'])
    assert card(page.abc_amount) is fno
    assert card(page.spins['astec_fno_percentual']) is fno
    assert card(page.spins['astec_of_percentual']) is of
    assert fno is not of
    assert isinstance(page.financing_shares['valor_fno'], QLabel)
    assert isinstance(page.financing_shares['valor_of'], QLabel)
    page.close()


@pytest.mark.parametrize('source', ['fno', 'of'])
def test_astec_toggling_no_preserves_entered_percentage_for_reopening(app, tmp_path, source):
    page = ProposalPage()
    proposal = Proposal()
    flag, percent = f'astec_{source}_financiada', f'astec_{source}_percentual'
    page.flags[flag].setValue(True)
    page.spins[percent].setValue(2.75)
    page.flags[flag].setValue(False)
    page.read_into(proposal)
    repository = SQLiteProposalRepository(tmp_path / 'inactive-choice.sqlite3')
    repository.save(proposal)
    reopened = repository.get(proposal.id)
    assert not getattr(reopened, flag) and getattr(reopened, percent) == Decimal('2.75')
    page.load(reopened)
    assert page.astec_fields[flag].isHidden()
    page.flags[flag].setValue(True)
    assert page.spins[percent].value() == 2.75
    assert not page.astec_fields[flag].isHidden()
    page.close()
    repository.engine.dispose()


@pytest.mark.parametrize('source', ['fno', 'of'])
def test_partial_draft_with_active_astec_and_unset_percentage_can_be_reviewed(app, tmp_path, source):
    proposal = Proposal(**{f'astec_{source}_financiada': True})
    proposal.validate()
    repository = SQLiteProposalRepository(tmp_path / 'partial.sqlite3')
    review = ReviewPage(ProposalService(repository, FakePropertyRepository()), AppSettings())
    review.load(proposal)
    assert f'% ASTEC {source.upper()}: —' in review.text.toPlainText()
    review.close()
    repository.engine.dispose()
