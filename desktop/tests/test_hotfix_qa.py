"""Manual-QA regressions using synthetic data and native Qt editing events."""
from decimal import Decimal
from hashlib import sha256
import sqlite3

import pytest
from PySide6.QtCore import Qt, QPointF, QPoint
from PySide6.QtGui import QWheelEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QWidget, QVBoxLayout, QLineEdit

from amazon_agro.domain.models import Proposal
from amazon_agro.config.settings import AppSettings
from amazon_agro.exporters.excel import OpenpyxlExcelProposalExporter
from amazon_agro.exporters.excel_map import EXCEL_FIELD_MAP
from amazon_agro.repositories.fake_properties import FakePropertyRepository
from amazon_agro.config.resources import resource_path
from openpyxl import load_workbook
from amazon_agro.integrations.workbook_inspector import NeedsConfigurationError
from amazon_agro.repositories.sqlite_proposals import SQLiteProposalRepository
from amazon_agro.ui.property_sources import PropertySourcesPage
from amazon_agro.ui.currency_input import CurrencyInput, parse_currency
from amazon_agro.ui.pages import OperationPage, ProposalPage
from test_property_workbooks import make_book, parse, record, catalog_sync


@pytest.fixture
def app(monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    return QApplication.instance() or QApplication([])


@pytest.mark.parametrize("text, value", [
    ("150000", "150000.00"), ("1.234,56", "1234.56"),
    ("1234,56", "1234.56"), ("R$ 1.234,56", "1234.56"),
    ("0,01", "0.01"), ("999999999999,99", "999999999999.99"),
])
def test_currency_parses_brazilian_amounts_exactly(text, value):
    assert parse_currency(text) == Decimal(value)


def test_currency_click_zero_typing_selection_partial_edit_and_paste(app):
    window = QWidget()
    layout = QVBoxLayout(window)
    money, next_field = CurrencyInput(), QLineEdit()
    layout.addWidget(money)
    layout.addWidget(next_field)
    window.show()
    next_field.setFocus()
    app.processEvents()
    QTest.mouseClick(money, Qt.MouseButton.LeftButton)
    QTest.keyClicks(money, "150000")
    assert money.value() == Decimal("150000.00")
    next_field.setFocus()
    app.processEvents()
    assert money.text() == "R$ 150.000,00"
    money.setFocus()
    QTest.keyClick(money, Qt.Key.Key_A, Qt.KeyboardModifier.ControlModifier)
    QTest.keyClicks(money, "1234,56")
    QTest.keyClick(money, Qt.Key.Key_Home)
    QTest.keyClick(money, Qt.Key.Key_Right)
    QTest.keyClick(money, Qt.Key.Key_Delete)
    assert money.value() == Decimal("134.56")
    QTest.keyClick(money, Qt.Key.Key_End)
    QTest.keyClick(money, Qt.Key.Key_Backspace)
    assert money.value() == Decimal("134.50")
    QTest.keyClick(money, Qt.Key.Key_A, Qt.KeyboardModifier.ControlModifier)
    app.clipboard().setText("R$ 1.234,56")
    QTest.keyClick(money, Qt.Key.Key_V, Qt.KeyboardModifier.ControlModifier)
    next_field.setFocus()
    app.processEvents()
    assert money.value() == Decimal("1234.56")
    assert money.text() == "R$ 1.234,56"
    window.close()


def test_currency_wheel_never_changes_value(app):
    money = CurrencyInput()
    money.setValue(Decimal("150000.01"))
    wheel = QWheelEvent(QPointF(5, 5), QPointF(5, 5), QPoint(), QPoint(0, 120),
                        Qt.MouseButton.NoButton, Qt.KeyboardModifier.NoModifier,
                        Qt.ScrollPhase.NoScrollPhase, False)
    QApplication.sendEvent(money, wheel)
    assert money.value() == Decimal("150000.01")
    money.close()


def test_financial_order_and_removed_operation_fields(app):
    operation = OperationPage(AppSettings())
    assert "status" not in operation.controls and "aguardando" not in operation.controls
    legacy = Proposal(status="Etapa legada", aguardando="Documento legado")
    operation.load(legacy)
    operation.read_into(legacy)
    assert legacy.status == "Etapa legada" and legacy.aguardando == "Documento legado"
    page = ProposalPage()
    assert list(page.spins) == ["valor_total", "valor_fno", "valor_of",
                              "astec_fno_percentual", "laudo_abc_percentual",
                              "astec_of_percentual", "percentual_recursos_proprios"]
    assert list(page.flags) == ["astec_fno_financiada", "astec_of_financiada"]
    assert all(isinstance(page.spins[key], CurrencyInput) for key in page.MONEY)
    operation.close()
    page.close()


def test_abc_exclusive_fields_persist_reopen_and_disable(app, tmp_path):
    page = ProposalPage()
    proposal = Proposal(astec_fno_percentual=Decimal("99"), astec_of_percentual=Decimal("77"))
    page.load(proposal)
    assert page.abc_no.isChecked() and not page.abc_yes.isChecked()
    assert page.abc_fields.isHidden() and not page.abc_fields.isEnabled()
    page.abc_yes.setChecked(True)
    assert not page.abc_no.isChecked()
    assert not page.abc_fields.isHidden() and page.abc_fields.isEnabled()
    page.spins["laudo_abc_percentual"].setValue(3.25)
    page.spins["valor_total"].setValue(Decimal("100000"))
    page.spins["valor_fno"].setValue(Decimal("70000"))
    page.read_into(proposal)
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    repository.save(proposal)
    repository.engine.dispose()
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    reopened = repository.get(proposal.id)
    assert reopened.laudo_abc_valor == Decimal("2275")
    assert reopened.laudo_abc_percentual == Decimal("3.25")
    assert reopened.astec_fno_percentual == Decimal("99")
    assert reopened.astec_of_percentual == Decimal("77")
    page.load(reopened)
    assert page.abc_amount.text() == "R$ 2.275,00"
    page.abc_no.setChecked(True)
    page.read_into(reopened)
    assert reopened.laudo_abc_percentual is reopened.laudo_abc_valor is None
    repository.save(reopened)
    assert not repository.get(proposal.id).laudo_abc_financiado
    repository.engine.dispose()
    page.close()


def test_abc_schema_migration_preserves_old_proposal_and_legacy_values(tmp_path):
    path = tmp_path / "legacy.sqlite3"
    repository = SQLiteProposalRepository(path)
    proposal = Proposal(status="Legado", aguardando="Legado", astec_fno_percentual=Decimal("2.75"))
    repository.save(proposal)
    repository.engine.dispose()
    with sqlite3.connect(path) as connection:
        connection.execute("ALTER TABLE proposals DROP COLUMN laudo_abc_valor")
    repository = SQLiteProposalRepository(path)
    reopened = repository.get(proposal.id)
    assert reopened.laudo_abc_valor is None
    assert reopened.status == reopened.aguardando == "Legado"
    assert reopened.astec_fno_percentual == Decimal("2.75")
    reopened.laudo_abc_financiado = True
    reopened.laudo_abc_valor = Decimal("1234.56")
    repository.save(reopened)
    assert repository.get(proposal.id).laudo_abc_valor == Decimal("1234.56")
    with sqlite3.connect(path) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
    repository.engine.dispose()


def test_basa_shared_document_merges_join_registration_blocks_without_filename_rules(tmp_path):
    # Anonymous reproduction: header row 7, B:J, separately merged names and
    # registrations, with CCIR/ITR/CAR physically shared between registrations.
    headers = ("Fazenda", "Matrículas", "Matrícula Anterior", "Lote/Gleba",
               "Proprietário", "CPF/CNPJ", "CCIR", "ITR", "CAR")
    rows, merges = [], []
    for index in range(5):
        start = len(rows) + 1
        rows.extend([{"Fazenda": "Fazenda Alfa" if index < 2 else "Fazenda Beta",
                      "Matrículas": f"MAT-{index + 1}", "Proprietário": "Pessoa Exemplo",
                      "CPF/CNPJ": "000.000.000-00",
                      "CCIR": "CCIR-A" if index == 0 else "CCIR-B" if index == 2 else None,
                      "ITR": "ITR-A" if index == 0 else "ITR-B" if index == 2 else None}, {}])
        merges.extend((field, start, start + 1) for field in
                      ("Fazenda", "Matrículas", "Proprietário", "CPF/CNPJ"))
    merges.extend((("CCIR", 1, 4), ("ITR", 1, 4), ("CCIR", 5, 10), ("ITR", 5, 10)))
    path = make_book(tmp_path / "nome-arbitrario.xlsx", rows, headers=headers,
                     header_row=7, merges=tuple(merges))
    before = sha256(path.read_bytes()).digest()
    parsed = parse(path)
    assert [len(f.parcels) for f in parsed.properties] == [2, 3]
    assert len(parsed.owners) == 1 and not parsed.warnings
    assert all(p.area is None for f in parsed.properties for p in f.parcels)
    assert sha256(path.read_bytes()).digest() == before


def test_shared_document_merge_does_not_join_different_farm_names(tmp_path):
    path = make_book(tmp_path / "different.xlsx", [
        {"Fazenda": "Fazenda Alfa", "Matrículas": "MAT-1", "CCIR": "CCIR-A"}, {},
        {"Fazenda": "Fazenda Beta", "Matrículas": "MAT-2"}, {},
    ], merges=(("Fazenda", 1, 2), ("Matrículas", 1, 2),
               ("Fazenda", 3, 4), ("Matrículas", 3, 4), ("CCIR", 1, 4)))
    assert len(parse(path).properties) == 2


def test_unassigned_basa_is_detected_by_headers_with_different_default_family(tmp_path):
    path = make_book(tmp_path / "arbitrary.xlsx", [record()], header_row=7)
    catalog, sync, settings = catalog_sync(tmp_path)
    settings.property_profile_name = "Outra família"
    settings.property_profile_options = {"column_aliases": {"name": ["Estabelecimento"]}}
    report = sync.synchronize(settings)
    assert report.events[0].status == "OK"
    assert catalog.get_source(path).profile == "BASA Ambiental"
    # An explicit manual override must never be replaced by automatic detection.
    settings.property_file_profiles[path.name] = "Outra família"
    assert sync.synchronize(settings).events[0].status == "NEEDS_CONFIGURATION"
    assert catalog.stats().properties == 1
    catalog.close()


def test_streaming_read_failure_preserves_catalog_and_explains_offline_retry(tmp_path, monkeypatch):
    path = make_book(tmp_path / "streaming.xlsx", [record()])
    catalog, sync, settings = catalog_sync(tmp_path)
    sync.synchronize(settings)
    before = catalog.search("")[0].external_id
    def unavailable(_path):
        raise OSError("Cloud provider unavailable")
    monkeypatch.setattr("amazon_agro.services.property_sync_service._fingerprint", unavailable)
    report = sync.synchronize(settings)
    assert report.events[0].status == "ERROR"
    assert catalog.search("")[0].external_id == before
    assert "disponível off-line" in catalog.get_source(path).last_error
    assert "Atualizar catálogo" in catalog.get_source(path).last_error
    catalog.close()


def test_source_screen_lists_discovered_names_and_visible_diagnostics(app, tmp_path):
    path = make_book(tmp_path / "not-yet-catalogued.xlsx", [record()])
    catalog, sync, settings = catalog_sync(tmp_path)
    page = PropertySourcesPage(settings, catalog, sync)
    assert page.found.text() == "Arquivos encontrados: 1"
    assert page.table.item(0, 0).text() == path.name
    assert page.table.item(0, 2).text() == "Aguardando atualização"
    sync.synchronize(settings)
    make_book(path, [{"Fazenda": "Sem matrícula"}])
    sync.synchronize(settings)
    page.refresh()
    assert page.table.item(0, 2).text() == "Precisa configurar"
    assert page.table.item(0, 3).text() == "1"
    assert "Nenhuma fazenda" in page.table.item(0, 4).text()
    page.close()
    catalog.close()


@pytest.mark.parametrize("enabled", [True, False])
def test_xlsx_abc_outputs_only_explicit_active_fields_and_ignores_legacy_percentages(tmp_path, enabled):
    proposal = Proposal(valor_total=Decimal("100000"), valor_fno=Decimal("70000"), laudo_abc_financiado=enabled, laudo_abc_percentual=Decimal("3.25"),
                        laudo_abc_valor=Decimal("4321.09"),
                        astec_fno_percentual=Decimal("77.88"), astec_of_percentual=Decimal("99.11"))
    path = OpenpyxlExcelProposalExporter(AppSettings(), FakePropertyRepository()).export(proposal, tmp_path / "abc.xlsx")
    sheet = load_workbook(path).active
    assert sheet[EXCEL_FIELD_MAP["laudo_abc_financiado"]].value == ("Sim" if enabled else "Não")
    assert sheet[EXCEL_FIELD_MAP["laudo_abc_percentual"]].value == (0.0325 if enabled else None)
    assert sheet[EXCEL_FIELD_MAP["laudo_abc_valor"]].value == (2275 if enabled else None)
    assert sheet[EXCEL_FIELD_MAP["astec_fno_percentual"]].value is None
    assert sheet[EXCEL_FIELD_MAP["astec_of_percentual"]].value is None
    assert not {0.7788, 0.9911}.intersection(cell.value for row in sheet for cell in row)
    assert "LOGO AMAZON" not in {cell.value for row in sheet for cell in row}
    assert len(sheet._images) == 1


def test_official_logo_resource_has_only_cropped_brand_pixels():
    from PIL import Image
    logo = resource_path("resources", "AmazonAgroLogo.png")
    with Image.open(logo) as image:
        assert image.size == (1221, 258)
        assert image.format == "PNG"


def test_official_amazon_document_migration_does_not_assign_it_to_another_consultancy():
    settings = AppSettings(default_consultancy_name="Outra Consultoria", default_consultancy_document="")
    assert settings.default_consultancy_document == ""


@pytest.mark.parametrize("field, invalid", [
    ("laudo_abc_valor", Decimal("-0.01")), ("laudo_abc_valor", Decimal("NaN")),
    ("laudo_abc_percentual", Decimal("100.01")), ("laudo_abc_percentual", Decimal("NaN")),
])
def test_active_abc_rejects_invalid_explicit_values(field, invalid):
    proposal = Proposal(laudo_abc_financiado=True)
    setattr(proposal, field, invalid)
    with pytest.raises(ValueError):
        proposal.validate()


@pytest.mark.parametrize("count, pages", [(1, 1), (4, 1), (5, 2), (8, 2), (9, 3)])
def test_logo_and_identity_repeat_on_each_xlsx_page(tmp_path, count, pages):
    from test_business_definitions import proposal_with_parcels
    path = OpenpyxlExcelProposalExporter(AppSettings(), FakePropertyRepository()).export(
        proposal_with_parcels(count), tmp_path / "pages.xlsx")
    sheet = load_workbook(path).active
    starts = [1] + [36 + index * 16 for index in range(pages - 1)]
    assert [image.anchor._from.row + 1 for image in sheet._images] == starts
    for start in starts:
        assert sheet[f"A{start + 2}"].value == "TEST-PAGINATION"
        assert sheet[f"A{start + 5}"].value == "Produtor Exemplo"
        if start > 1:
            values = {cell.value for row in sheet.iter_rows(min_row=start, max_row=start+15) for cell in row}
            assert "III - PROPOSTA" not in values and "VALOR LAUDO ABC" not in values
