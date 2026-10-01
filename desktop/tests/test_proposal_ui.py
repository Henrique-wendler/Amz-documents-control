"""Synthetic drafts only. Exercise Qt clicks through SQLite and the real XLSX exporter."""
from decimal import Decimal
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import Qt, QTimer
from PySide6.QtGui import QDesktopServices
from PySide6.QtTest import QTest, QSignalSpy
from PySide6.QtWidgets import QApplication, QDialog, QFileDialog, QMessageBox, QWizard

from amazon_agro.config.settings import AppSettings, default_data_dir
from amazon_agro.domain.models import (
    Participant, ParticipantType, Proposal, ProposalProperty, ProposalPropertyParcel,
    PropertyClassification,
)
from amazon_agro.exporters.excel import OpenpyxlExcelProposalExporter
from amazon_agro.exporters.pdf import SpreadsheetPdfProposalExporter
from amazon_agro.exporters.pdf_backends import PdfBackendSelector
from amazon_agro.repositories.sqlite_proposals import SQLiteProposalRepository
from amazon_agro.repositories.sqlite_property_catalog import SQLitePropertyCatalogRepository
from amazon_agro.services.export_service import ProposalExportService
from amazon_agro.services.export_validator import ProposalExportValidator
from amazon_agro.services.proposal_service import ProposalService
from amazon_agro.services.property_sync_service import PropertyCatalogSyncService
from amazon_agro.ui.export_dialogs import ExportSuccessDialog, offer_excel
from amazon_agro.ui.main_window import MainWindow
from amazon_agro.ui.onboarding import FirstRunWizard
from amazon_agro.ui.settings_dialog import SettingsDialog
from amazon_agro.ui.conversion_worker import run_conversion


@pytest.fixture
def ui(tmp_path, monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    settings = AppSettings.load(tmp_path / "config.json")
    settings.default_output_dir = str(tmp_path)
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    catalog = SQLitePropertyCatalogRepository(tmp_path / "catalog.sqlite3")
    validator = ProposalExportValidator(settings, catalog)
    export = ProposalExportService(repository, validator,
        OpenpyxlExcelProposalExporter(settings, catalog),
        SpreadsheetPdfProposalExporter(PdfBackendSelector([])))
    window = MainWindow(ProposalService(repository, catalog), settings, export,
                        validator, catalog, PropertyCatalogSyncService(catalog))
    window.show()
    yield window, app
    window.refresh_timer.stop()
    window.close()
    app.processEvents()
    repository.engine.dispose()
    catalog.close()


def fill(window):
    window.new_proposal()
    for field, value in {
        "numero_proposta": "TEST-004", "proponente": "Proponente Exemplo",
        "cpf_cnpj": "000.000.000-00", "tecnico": "Técnico Exemplo",
        "finalidade": "Custeio", "agencia": "Agência Exemplo", "fonte": "FNO",
    }.items():
        control = window.operation_page.controls[field]
        if hasattr(control, "setCurrentText"):
            control.setCurrentText(value)
        else:
            control.setText(value)
    window.proposal_page.spins["valor_total"].setValue(123456.78)


def click_review(window, app, kind):
    window.steps.setCurrentRow(4)
    app.processEvents()
    button = {"xlsx": window.review_page.excel_button,
              "pdf": window.review_page.pdf_button,
              "both": window.review_page.both_button}[kind]
    spy = QSignalSpy(button.clicked)
    QTest.mouseClick(button, Qt.MouseButton.LeftButton)
    assert spy.count() == 1


def test_initial_state_has_no_draft(ui):
    window, _ = ui
    assert not window.has_proposal
    assert window.workspace.currentIndex() == 0
    assert not window.header_buttons["Salvar"].isEnabled()
    assert window.steps.count() == 5


def test_summary_live_without_saving_and_uses_validator(ui):
    window, app = ui
    fill(window)
    app.processEvents()
    text = window.summary.details.text()
    for value in ("TEST-004", "Proponente Exemplo", "Técnico Exemplo", "FNO", "123.456,78"):
        assert value in text
    assert "000.000.000-00" not in text
    assert "pronta para gerar" in window.summary.status.text()
    assert not window.service.list_recent()
    window.operation_page.controls["finalidade"].clear()
    app.processEvents()
    errors = window.export_validator.validate(window._collect()).errors
    assert all(error in window.summary.pending.text() for error in errors)
    assert "Informe a finalidade." in window.summary.pending.text()


def test_summary_tracks_participants_properties_and_removal(ui):
    window, app = ui
    fill(window)
    window.participants_page.add_participant(Participant(window.current.id, "Pessoa Exemplo"))
    window.properties_page._add_link(ProposalProperty(
        window.current.id, "synthetic-farm", PropertyClassification.CREDIT_OBJECT,
        property_name_snapshot="Fazenda Exemplo", municipality_snapshot="Palmas",
        state_snapshot="TO", selected_parcels=[ProposalPropertyParcel("parcel-1", "MAT-004")],
    ))
    app.processEvents()
    assert "Classificação de imóvel inválida." in window.summary.pending.text()
    # Review must remain reachable even while classification is missing.
    window.steps.setCurrentRow(4)
    assert "Classificação pendente" in window.review_page.text.toPlainText()
    window.properties_page.selected.cellWidget(0, 4).setCurrentIndex(2)
    window.participants_page.table.cellWidget(0, 2).setCurrentIndex(3)
    app.processEvents()
    for value in ("Pessoa Exemplo", "Avalista", "Fazenda Exemplo", "Palmas", "MAT-004"):
        assert value in window.summary.details.text()
    assert "pronta para gerar" in window.summary.status.text()
    window.properties_page.selected.setCurrentCell(0, 0)
    window.properties_page.remove_selected()
    window.participants_page.remove_selected()
    app.processEvents()
    assert "Participantes (0)" in window.summary.details.text()
    assert "Imóveis (0)" in window.summary.details.text()


@pytest.mark.parametrize("kind", ["xlsx", "pdf", "both"])
def test_pending_export_click_explains_instead_of_disabling(ui, monkeypatch, kind):
    window, app = ui
    window.new_proposal()
    messages = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: messages.append(args[2]))
    click_review(window, app, kind)
    assert len(messages) == 1
    assert "Campos pendentes" in messages[0]
    assert "técnico" in messages[0]
    assert not window.service.list_recent()


def test_click_xlsx_saves_and_produces_real_workbook_without_pdf(ui, monkeypatch, tmp_path):
    window, app = ui
    fill(window)
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *args: str(tmp_path))
    results = []
    monkeypatch.setattr(ExportSuccessDialog, "exec", lambda dialog: results.extend(dialog.paths))
    click_review(window, app, "xlsx")
    assert len(results) == 1 and results[0].suffix == ".xlsx"
    workbook = load_workbook(results[0])
    assert workbook.active["A6"].value == "Proponente Exemplo"
    assert workbook.active["K3"].value == "TEST-004"
    workbook.close()
    assert window.service.get(window.current.id).valor_total == Decimal("123456.78")


@pytest.mark.parametrize("kind", ["pdf", "both"])
def test_pdf_unavailable_offers_working_excel(ui, monkeypatch, tmp_path, kind):
    window, app = ui
    fill(window)
    offered = []
    monkeypatch.setattr("amazon_agro.ui.main_window.offer_excel", lambda parent: offered.append(True) or True)
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *args: str(tmp_path))
    results = []
    monkeypatch.setattr(ExportSuccessDialog, "exec", lambda dialog: results.extend(dialog.paths))
    click_review(window, app, kind)
    assert offered == [True]
    assert len(results) == 1 and results[0].is_file()
    assert results[0].suffix == ".xlsx"


@pytest.mark.parametrize("kind", ["pdf", "both"])
def test_click_pdf_routes_saved_proposal_and_results(ui, monkeypatch, tmp_path, kind):
    window, app = ui
    fill(window)
    calls = []
    class Backend:
        name = "Synthetic PDF backend"
        def is_available(self):
            return True
        def convert(self, source, destination):
            assert window.service.get(window.current.id) is not None
            assert load_workbook(source).active["A6"].value == "Proponente Exemplo"
            calls.append(source)
            destination.write_bytes(b"%PDF-1.4\nsynthetic conversion test")
    window.export_service.pdf.selector = PdfBackendSelector([Backend()])
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *args: str(tmp_path))
    results = []
    monkeypatch.setattr(ExportSuccessDialog, "exec", lambda dialog: results.extend(dialog.paths))
    click_review(window, app, kind)
    assert len(calls) == 1
    assert {p.suffix for p in results} == ({".pdf"} if kind == "pdf" else {".pdf", ".xlsx"})


@pytest.mark.parametrize("error, expected", [
    (PermissionError("technical private detail"), "permissão"),
    (FileNotFoundError("technical private detail"), "não foi encontrado"),
    (FileExistsError("technical private detail"), "Já existe"),
    (RuntimeError("technical private detail"), "detalhes técnicos"),
])
def test_export_error_visible_and_details_only_in_log(ui, monkeypatch, tmp_path, caplog, error, expected):
    window, app = ui
    fill(window)
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *args: str(tmp_path))
    def fail(*args):
        raise error
    monkeypatch.setattr(window.export_service, "generate_excel", fail)
    messages = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: messages.append(args[2]))
    click_review(window, app, "xlsx")
    assert expected in messages[0]
    assert "technical private detail" not in messages[0]
    assert "technical private detail" in caplog.text
    assert window.workspace.isEnabled()
    assert QApplication.overrideCursor() is None


def test_export_cancel_has_feedback(ui, monkeypatch):
    window, app = ui
    fill(window)
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *args: "")
    click_review(window, app, "xlsx")
    assert "cancelada" in window.statusBar().currentMessage()
    assert not window.service.list_recent()


def test_success_dialog_opens_selected_file_and_folder(ui, tmp_path, monkeypatch):
    paths = [tmp_path / "test.xlsx", tmp_path / "test.pdf"]
    for path in paths:
        path.write_bytes(b"synthetic")
    dialog = ExportSuccessDialog(paths)
    dialog.show()
    opened = []
    monkeypatch.setattr(QDesktopServices, "openUrl", lambda url: opened.append(url.toLocalFile()) or True)
    dialog.file_choice.setCurrentIndex(1)
    QTest.mouseClick(dialog.open_file_button, Qt.MouseButton.LeftButton)
    QTest.mouseClick(dialog.open_folder_button, Qt.MouseButton.LeftButton)
    assert [Path(p) for p in opened] == [paths[1], tmp_path]
    dialog.close()


def test_open_failure_is_visible(ui, tmp_path, monkeypatch):
    path = tmp_path / "example.xlsx"
    path.write_bytes(b"synthetic")
    dialog = ExportSuccessDialog([path])
    messages = []
    monkeypatch.setattr(QDesktopServices, "openUrl", lambda url: False)
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: messages.append(args[2]))
    dialog.open_file()
    assert "Explorador" in messages[0]
    dialog.close()


def test_unavailable_dialog_real_button(ui):
    window, app = ui
    def choose():
        dialog = app.activeModalWidget()
        assert "indisponível" in dialog.windowTitle()
        button = next(b for b in dialog.buttons() if b.text() == "Gerar Excel")
        QTest.mouseClick(button, Qt.MouseButton.LeftButton)
    QTimer.singleShot(0, choose)
    assert offer_excel(window)


def test_first_run_wizard_saves_once_and_defaults(ui, tmp_path, monkeypatch):
    window, app = ui
    calls = []
    def complete(wizard):
        calls.append(True)
        wizard.source.setText(str(tmp_path))
        wizard.output.setText(str(tmp_path / "exports"))
        wizard.city.setText("Cidade Exemplo")
        wizard.technician.setText("Técnico Exemplo")
        wizard.accept()
        return wizard.result()
    monkeypatch.setattr(FirstRunWizard, "exec", complete)
    window.startup()
    loaded = AppSettings.load(tmp_path / "config.json")
    assert loaded.first_run_completed
    assert loaded.source_directory() == tmp_path
    assert loaded.output_dir() == tmp_path / "exports"
    assert loaded.output_dir().is_dir()
    assert window.current.cidade == "Cidade Exemplo"
    assert window.current.tecnico == "Técnico Exemplo"
    window.settings = loaded
    window.startup()
    assert len(calls) == 1


def test_repeat_wizard_from_general_settings(ui, monkeypatch):
    window, app = ui
    window.settings.first_run_completed = True
    calls = []
    monkeypatch.setattr(FirstRunWizard, "exec", lambda wizard: calls.append(True) or QDialog.DialogCode.Rejected)
    dialog = SettingsDialog(window.settings, window.catalog, window.sync_service)
    dialog.show()
    QTest.mouseClick(dialog.repeat, Qt.MouseButton.LeftButton)
    assert calls == [True]
    assert window.settings.first_run_completed
    dialog.close()


def test_cancel_wizard_does_not_commit(ui, tmp_path):
    window, _ = ui
    wizard = FirstRunWizard(window.settings)
    wizard.city.setText("Changed")
    wizard.reject()
    assert not window.settings.first_run_completed
    assert window.settings.default_city != "Changed"
    assert not (tmp_path / "config.json").exists()


def test_wizard_folder_browse_and_skip(ui, tmp_path, monkeypatch):
    window, _ = ui
    wizard = FirstRunWizard(window.settings)
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *args: str(tmp_path))
    wizard.source.browse()
    assert wizard.source.text() == str(tmp_path)
    wizard.show()
    wizard.next()
    wizard._skip_source()
    assert wizard.currentId() == 2
    assert wizard.source.text() == ""
    wizard.reject()


@pytest.mark.parametrize("available", [[], ["Excel COM"], ["LibreOffice"], ["Excel COM", "LibreOffice"]])
def test_wizard_pdf_detection(ui, available):
    class Backend:
        def __init__(self, name): self.name = name
        def is_available(self): return self.name in available
    window, _ = ui
    wizard = FirstRunWizard(window.settings, selector=PdfBackendSelector([
        Backend("Excel COM"), Backend("LibreOffice")]))
    wizard._page_changed(3)
    assert ("PDF disponível" if available else "PDF indisponível") in wizard.pdf_status.text()
    assert wizard.pdf_status.text().count("✓ Encontrado") == len(available)
    wizard.reject()


@pytest.mark.parametrize("size", [(1366, 768), (1920, 1080)])
def test_window_layout_and_collapsible_summary(ui, size):
    window, app = ui
    fill(window)
    window.resize(*size)
    app.processEvents()
    assert window.width() == size[0]
    for index in range(5):
        window.steps.setCurrentRow(index)
        app.processEvents()
        scroll = window.stack.widget(index)
        assert not scroll.horizontalScrollBar().isVisible()
        assert scroll.widget().minimumSizeHint().width() <= scroll.viewport().width()
    window.toggle_summary()
    assert window.summary_scroll.isHidden()
    window.toggle_summary()
    assert not window.summary_scroll.isHidden()
    assert window.splitter.sizes()[2] > 0


def test_resource_paths_independent_of_cwd_and_frozen(tmp_path, monkeypatch):
    settings = AppSettings.load(tmp_path / "config.json")
    template = settings.template_path()
    monkeypatch.chdir(tmp_path)
    assert settings.template_path() == template
    settings.xlsx_template_path = "models/custom.xlsx"
    assert settings.template_path() == tmp_path / "models/custom.xlsx"
    settings.xlsx_template_path = ""
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "_MEIPASS", str(tmp_path / "bundle"), raising=False)
    assert settings.template_path() == tmp_path / "bundle/amazon_agro/templates/modelo_proposta.xlsx"
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    assert default_data_dir() == tmp_path / "AmazonAgro"


def test_conversion_keeps_qt_event_loop_running(ui, tmp_path):
    from threading import Event
    event = Event()
    class Backend:
        def convert(self, source, destination):
            assert event.wait(2), "The GUI event loop did not run while conversion was active"
            destination.write_bytes(b"%PDF-synthetic")
            return destination
    QTimer.singleShot(0, event.set)
    path = tmp_path / "converted.pdf"
    assert run_conversion(Backend(), tmp_path / "source.xlsx", path) == path
    assert path.is_file()


def test_conversion_worker_propagates_failure(ui, tmp_path):
    class Backend:
        def convert(self, *args):
            raise RuntimeError("conversion failed")
    with pytest.raises(RuntimeError, match="conversion failed"):
        run_conversion(Backend(), tmp_path / "source", tmp_path / "target")


def test_busy_conversion_protects_draft_and_close(ui):
    window, app = ui
    fill(window)
    original_id = window.current.id
    window._exporting = True
    window.new_proposal()
    window.close()
    assert window.current.id == original_id
    assert window.isVisible()
    assert "Aguarde" in window.statusBar().currentMessage()
    window._exporting = False


def test_wizard_save_failure_does_not_mark_complete(ui, monkeypatch):
    window, _ = ui
    wizard = FirstRunWizard(window.settings)
    def fail(*args):
        raise PermissionError("configuration read only")
    monkeypatch.setattr(AppSettings, "save", fail)
    messages = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: messages.append(args[2]))
    wizard.accept()
    assert not window.settings.first_run_completed
    assert "permissões" in messages[0]
    wizard.reject()


def test_real_wizard_navigation_and_finish(ui, tmp_path):
    window, app = ui
    wizard = FirstRunWizard(window.settings, selector=PdfBackendSelector([]))
    wizard.output.setText(str(tmp_path / "wizard-output"))
    wizard.show()
    app.processEvents()
    for index in range(5):
        assert wizard.currentId() == index
        QTest.mouseClick(wizard.button(QWizard.WizardButton.NextButton), Qt.MouseButton.LeftButton)
        app.processEvents()
    QTest.mouseClick(wizard.button(QWizard.WizardButton.FinishButton), Qt.MouseButton.LeftButton)
    assert wizard.result() == QDialog.DialogCode.Accepted
    assert AppSettings.load(tmp_path / "config.json").first_run_completed


def test_navigation_marks_unvisited_pending_and_complete(ui):
    window, app = ui
    window.new_proposal()
    assert "Não visitada" in window.steps.item(2).text()
    window.steps.setCurrentRow(2)
    assert "Com pendência" in window.steps.item(0).text()
    window.steps.setCurrentRow(1)
    assert "Com pendência" in window.steps.item(2).text()
    window.steps.setCurrentRow(4)
    assert "Completa" in window.steps.item(1).text()
