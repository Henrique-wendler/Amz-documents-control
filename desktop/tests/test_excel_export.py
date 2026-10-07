from copy import copy
from datetime import date
from decimal import Decimal
from hashlib import sha256

import pytest
from openpyxl import load_workbook

from amazon_agro.config.settings import AppSettings
from amazon_agro.domain.models import (
    Participant, ParticipantType, Proposal, ProposalProperty,
    PropertyClassification,
)
from amazon_agro.exporters.excel import OpenpyxlExcelProposalExporter
from amazon_agro.exporters.excel_map import EXCEL_FIELD_MAP, PDF_PRINT_AREA
from amazon_agro.repositories.fake_properties import FakePropertyRepository


def filled_proposal() -> Proposal:
    proposal = Proposal(
        numero_proposta="12345", agencia="Palmas Centro",
        proponente="João da Silva", cpf_cnpj="123.456.789-00",
        tecnico="Maria Técnica", porte="Pequeno", finalidade="Custeio",
        descricao="Cultivo de cacau", fonte="FNO", cidade="Palmas",
        data_proposta=date(2026, 9, 29),
        recursos_proprios=Decimal("100000.00"),
        percentual_recursos_proprios=Decimal("12.50"),
        valor_total=Decimal("1234567.89"),
        valor_fno=Decimal("900000.00"),
        classificacao_da_percentual=Decimal("7"),
        astec_fno_financiada=True,
        astec_fno_percentual=Decimal("2.50"),
        laudo_abc_financiado=False,
        valor_of=Decimal("234567.89"),
        astec_of_financiada=True,
        astec_of_percentual=Decimal("1.25"),
    )
    proposal.participants = [
        Participant(
            proposal.id, "João da Silva", "123.456.789-00",
            ParticipantType.MAIN_ISSUER,
        ),
        Participant(
            proposal.id, "Ana Souza", "987.654.321-00",
            ParticipantType.GUARANTOR,
        ),
    ]
    proposal.properties = [
        ProposalProperty(proposal.id, "DEMO-001", PropertyClassification.CLASS_1),
        ProposalProperty(
            proposal.id, "DEMO-002", PropertyClassification.CREDIT_OBJECT,
        ),
    ]
    return proposal


def test_xlsx_integration_fills_mapped_cells_and_preserves_template(tmp_path) -> None:
    settings = AppSettings()
    template = settings.template_path()
    original_bytes = template.read_bytes()
    source = load_workbook(template)
    proposal = filled_proposal()
    destination = tmp_path / "output.xlsx"

    OpenpyxlExcelProposalExporter(
        settings, FakePropertyRepository()
    ).export(proposal, destination)

    assert sha256(template.read_bytes()).digest() == sha256(original_bytes).digest()
    output = load_workbook(destination)
    source_sheet, sheet = source.active, output.active
    assert output.sheetnames == source.sheetnames
    assert sheet["A3"].value == "12345"
    assert sheet["A2"].value == "Nº DA PROPOSTA"
    assert sheet["K3"].value == "Palmas Centro"
    assert sheet["K2"].value == "AGÊNCIA"
    assert sheet["A6"].value == "João da Silva"
    assert sheet["D6"].value == "123.456.789-00"
    assert sheet["I6"].value == "Maria Técnica"
    assert sheet["F6"].value == "1 participante(s)"
    assert sheet["A9"].value == "João da Silva"
    assert sheet["I10"].value == "987.654.321-00"
    assert sheet["M10"].value == 4
    assert sheet["A15"].value is None
    assert sheet["I15"].value is None
    assert sheet["M15"].value is None
    assert sheet["A21"].value == "Custeio"
    assert sheet["C21"].value == "Cultivo de cacau"
    assert sheet["G21"].value == "Sim"
    assert sheet["H21"].value == pytest.approx(0.125)
    assert sheet["H21"].number_format == "0.00%"
    assert sheet["I21"].value == "FNO"
    assert sheet["K21"].value == 1234567.89
    assert sheet["A23"].value == 900000
    assert sheet["K21"].number_format == '"R$" #,##0.00'
    assert sheet["D23"].value == pytest.approx(float(proposal.valor_fno / proposal.valor_total))
    assert sheet["D25"].value == pytest.approx(float(proposal.valor_of / proposal.valor_total))
    assert sheet["D23"].number_format == sheet["D25"].number_format == "0.00%"
    assert "classificacao_da_percentual" not in EXCEL_FIELD_MAP
    assert sheet["F23"].value == "Sim"
    assert sheet["I23"].value == "Não"
    assert sheet["K23"].value is None
    assert sheet["A25"].value == 234567.89
    assert sheet["F25"].value == "Sim"
    assert sheet["H23"].value == 0.025 and sheet["H25"].value == 0.0125
    assert sheet["H23"].number_format == sheet["H25"].number_format == "0.00%"
    assert sheet["I25"].value is None
    assert sheet["A28"].value == "Imóvel demonstrativo A"
    assert sheet["E28"].value == "Palmas"
    assert sheet["I28"].value == "MAT-001"
    assert sheet["K28"].value == 1
    assert sheet["K29"].value == 2
    assert sheet["A32"].value == "1 - HIPOTECA"
    assert sheet["A33"].value is None
    assert sheet["A34"].value == "TÉCNICO RESPONSÁVEL: Maria Técnica"
    assert sheet["A35"].value == "Palmas, 29 de setembro de 2026"
    assert "recursos_proprios" not in EXCEL_FIELD_MAP
    assert {cell.value for row in sheet.iter_rows(min_row=33, max_row=35, max_col=13)
            for cell in row if cell.value is not None} == {
        "TÉCNICO RESPONSÁVEL: Maria Técnica",
        "Palmas, 29 de setembro de 2026",
    }
    assert all(cell.value is None for row in sheet.iter_rows(min_row=36)
               for cell in row)
    assert EXCEL_FIELD_MAP["valor_total"] == "K21"

    source_merges = {str(item) for item in source_sheet.merged_cells.ranges}
    assert {str(item) for item in sheet.merged_cells.ranges} == (
        source_merges - {"A2:B3", "F22:H22", "F24:H24", "F25:H25",
                         "J22:L22", "J23:L23", "J24:L24", "J25:L25"}
    ) | {"A2:B2", "A3:B3", "C21:F21", "F22:G22", "F23:G23",
         "F24:G24", "F25:G25", "I22:J22", "K22:M22",
         "I23:J23", "K23:M23", "I24:M24", "I25:M25"}
    # The generated copy adopts the approved Word grid. Its original template
    # bytes remain identical, while text gets readable widths and row heights.
    assert sum(dim.width for dim in sheet.column_dimensions.values()) < 100
    assert 750 < sum(sheet.row_dimensions[row].height for row in range(1, 36)) < 850
    for coordinate in ("A2", "A6", "A15", "A21", "K21", "A28", "K28", "A32"):
        after = sheet[coordinate]
        assert after.font.name == "Times New Roman" and after.font.sz >= 8.5
        assert after.border.left.style == after.border.right.style == "thin"
        assert after.alignment.vertical == "center"
        assert after.alignment.wrap_text or after.alignment.shrinkToFit
    assert len(sheet._images) == 1
    assert sheet._images[0].anchor._from.row == 0
    assert sheet.page_setup.orientation == source_sheet.page_setup.orientation
    assert sheet.page_setup.scale is None
    assert sheet.print_area == source_sheet.print_area
    assert "$A$1:$M$35" in str(sheet.print_area)
    assert PDF_PRINT_AREA == "A1:M35"
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert sheet.page_setup.fitToWidth == sheet.page_setup.fitToHeight == 1
    assert len(sheet.row_breaks.brk) == 0
    assert sheet.page_margins.left >= 0.3 and sheet.page_margins.right >= 0.3


def test_xlsx_rejects_more_than_seven_participants(tmp_path) -> None:
    proposal = filled_proposal()
    proposal.participants = [
        Participant(proposal.id, f"Pessoa {index}") for index in range(8)
    ]
    with pytest.raises(ValueError, match="até 7 participantes"):
        OpenpyxlExcelProposalExporter(
            AppSettings(), FakePropertyRepository()
        ).export(proposal, tmp_path / "overflow.xlsx")


def test_xlsx_keeps_recursos_proprios_as_indicator_without_monetary_footer(tmp_path) -> None:
    proposal = filled_proposal()
    proposal.recursos_proprios = Decimal("0")
    proposal.percentual_recursos_proprios = Decimal("0")
    destination = tmp_path / "no-own-funds.xlsx"

    OpenpyxlExcelProposalExporter(
        AppSettings(), FakePropertyRepository()
    ).export(proposal, destination)

    sheet = load_workbook(destination).active
    assert sheet["G21"].value == "Não"
    assert sheet["H21"].value == 0
    assert sheet["A34"].value == "TÉCNICO RESPONSÁVEL: Maria Técnica"
    assert sheet["A35"].value == "Palmas, 29 de setembro de 2026"
    assert all(
        "RECURSOS PRÓPRIOS" not in str(cell.value)
        for row in sheet.iter_rows(min_row=33, max_row=35, max_col=13)
        for cell in row if cell.value is not None
    )


def test_xlsx_paginates_more_than_four_properties(tmp_path) -> None:
    from amazon_agro.domain.models import ProposalPropertyParcel
    proposal = filled_proposal()
    proposal.properties = [
        ProposalProperty(proposal.id, f"X-{index}", PropertyClassification.CLASS_1,
                         property_name_snapshot=f"Farm {index}",
                         selected_parcels=[ProposalPropertyParcel(f"P-{index}", f"MAT-{index}")])
        for index in range(5)
    ]
    path = OpenpyxlExcelProposalExporter(
        AppSettings(), FakePropertyRepository()
    ).export(proposal, tmp_path / "overflow.xlsx")
    sheet = load_workbook(path).active
    assert sheet["I44"].value == "MAT-4"
    assert len(sheet.row_breaks.brk) == 1




def test_xlsx_rejects_wrong_template_structure(tmp_path) -> None:
    from amazon_agro.exporters.excel import TemplateMappingError

    workbook = load_workbook(AppSettings().template_path())
    workbook.active["A7"] = "Outro modelo"
    wrong_template = tmp_path / "wrong.xlsx"
    workbook.save(wrong_template)
    settings = AppSettings(xlsx_template_path=str(wrong_template))
    with pytest.raises(TemplateMappingError, match="A7"):
        OpenpyxlExcelProposalExporter(
            settings, FakePropertyRepository()
        ).export(filled_proposal(), tmp_path / "out.xlsx")
