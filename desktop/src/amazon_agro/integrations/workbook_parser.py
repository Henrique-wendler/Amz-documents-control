"""Profile-driven, read-only parsing of property and registration hierarchy."""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from amazon_agro.domain.models import ParcelOwner, PropertyOwner, PropertyParcel, RuralProperty
from amazon_agro.domain.owner_identity import normalized_owner_document, owner_id
from amazon_agro.integrations.area_parser import parse_area_ha
from amazon_agro.integrations.workbook_inspector import NeedsConfigurationError, select_header
from amazon_agro.integrations.workbook_profile import PropertyWorkbookProfile


PARSER_VERSION = 4
_PARCEL_FIELDS = ("area", "previous_registration", "lot_description", "ccir", "itr", "car")
_LABELS = {
    "area": "Área", "previous_registration": "Matrícula anterior",
    "lot_description": "Lote/Gleba", "ccir": "CCIR", "itr": "ITR", "car": "CAR",
    "municipality": "Município", "state": "UF",
}


@dataclass(frozen=True, slots=True)
class ParsedWorkbook:
    source_file: str
    profile_name: str
    owner_name: str
    owner_document: str
    properties: tuple[RuralProperty, ...]
    warnings: tuple[str, ...] = ()

    @property
    def owners(self) -> tuple[PropertyOwner, ...]:
        return tuple({owner.id: owner for item in self.properties for owner in item.owners}.values())


@dataclass(frozen=True)
class _OwnerObservation:
    name: str
    document: str
    name_cell: str
    document_cell: str


@dataclass
class _ParcelRecord:
    row: int
    name: str
    registration: str
    group_key: tuple[str, int]
    values: dict[str, str] = field(default_factory=dict)
    owners: set[tuple[str, str]] = field(default_factory=set)
    last_row: int = 0
    owner_observations: list[_OwnerObservation] = field(default_factory=list)


def _text(value: object, number_format: str = "") -> str:
    if value is None:
        return ""
    if isinstance(value, int) and re.fullmatch(r"0+", number_format):
        return str(value).zfill(len(number_format))
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def _id(prefix: str, fields: tuple[str, ...], values: dict[str, str]) -> str:
    parts = []
    for key in fields:
        value = values.get(key, "")
        if key == "owner_document":
            digits = re.sub(r"\D", "", value)
            value = digits or value
        parts.append(" ".join(value.casefold().split()))
    if not any(parts):
        raise NeedsConfigurationError("Campos de identidade vazios no perfil.")
    payload = json.dumps(parts, ensure_ascii=False, separators=(",", ":"))
    return prefix + hashlib.sha256(payload.encode("utf-8")).hexdigest()[:32]


def _legacy_identity_owners(owners: set[tuple[str, str]]) -> set[tuple[str, str]]:
    """Reproduce old property ID inputs, never normalized owner membership."""
    canonical = {}
    for name, document in sorted(owners):
        key = (name.casefold(), re.sub(r"\D", "", document) or document.casefold())
        canonical.setdefault(key, (name, document))
    return {
        value for key, value in canonical.items()
        if not any(other != key and all(not part or part == other[index]
                                       for index, part in enumerate(key))
                   for other in canonical)
    }


class PropertyWorkbookParser:
    @staticmethod
    def _vertical_merges(sheet: Worksheet) -> dict[tuple[int, int], tuple[int, int]]:
        result = {}
        for merged in sheet.merged_cells.ranges:
            if merged.min_col != merged.max_col or merged.min_row == merged.max_row:
                continue
            for row in range(merged.min_row, merged.max_row + 1):
                result[row, merged.min_col] = (merged.min_row, merged.min_col)
        return result

    @staticmethod
    def _read(
        sheet: Worksheet, row: int, column: int | None,
        merges: dict[tuple[int, int], tuple[int, int]], inherit_merged: bool,
    ) -> str:
        if column is None:
            return ""
        cell = sheet.cell(row, column)
        if cell.value is None and inherit_merged:
            cell = sheet.cell(*merges.get((row, column), (row, column)))
        return _text(cell.value, cell.number_format)

    def parse(self, path: Path, relative_path: str,
              profile: PropertyWorkbookProfile) -> ParsedWorkbook:
        workbook = load_workbook(path, data_only=False, read_only=False)
        try:
            sheet, header = select_header(workbook, profile)
            columns = header.columns
            merges = self._vertical_merges(sheet)
            warnings: list[str] = []
            errors: list[str] = []
            records: dict[int, _ParcelRecord] = {}
            registration_column = columns["registration"]

            def read(row: int, key: str, inherit: bool = True) -> str:
                return self._read(sheet, row, columns.get(key), merges, inherit)

            def registration_anchor(row: int) -> int:
                return merges.get((row, registration_column), (row, registration_column))[0]

            def owner_observation(row: int) -> _OwnerObservation:
                def value_and_cell(key: str, fallback: str) -> tuple[str, str]:
                    value = read(row, key)
                    column = columns.get(key)
                    coordinate = sheet.cell(*merges.get((row, column), (row, column))).coordinate if column else ""
                    if not value and fallback:
                        return _text(sheet[fallback].value), sheet[fallback].coordinate
                    return value, coordinate
                name, name_cell = value_and_cell("owner_name", profile.owner_name_cell)
                document, document_cell = value_and_cell("owner_document", profile.owner_document_cell)
                return _OwnerObservation(name, document, name_cell, document_cell)

            def merged_targets(row: int, key: str) -> set[int]:
                """Registration anchors explicitly covered by this field's merge."""
                column = columns.get(key)
                anchor = merges.get((row, column))
                if anchor is None:
                    return set()
                return {
                    registration_anchor(other_row)
                    for (other_row, other_column), other_anchor in merges.items()
                    if other_column == column and other_anchor == anchor
                    and read(other_row, "registration")
                }

            active_name = ""
            active_group = 0
            previous_name = ""
            property_rows: list[int] = []
            for row in range(header.row + 1, sheet.max_row + 1):
                direct = {key: read(row, key, False) for key in columns}
                outside = [cell.coordinate for cell in sheet[row]
                           if cell.column not in columns.values() and _text(cell.value)]
                if outside:
                    warnings.append(
                        f"Linha {row}: conteúdo fora das colunas da tabela ignorado "
                        f"({', '.join(outside)})."
                    )
                formula_fields = [key for key, column in columns.items()
                                  if sheet.cell(row, column).data_type in {"f", "e"}]
                if formula_fields:
                    errors.append(f"Linha {row}: fórmula ou erro Excel em campo mapeado; "
                                  "forneça um valor verificável para importação.")
                    continue
                name = read(row, "name")
                registration = read(row, "registration")
                if not any(direct.values()) and not registration:
                    active_name = ""
                    previous_name = ""
                    continue
                if not name and profile.inherit_unmerged_property_name:
                    name = active_name
                if name:
                    active_name = name
                name_anchor = merges.get((row, columns["name"]))
                if name_anchor is not None:
                    group_key = (name, name_anchor[0])
                else:
                    if name != previous_name:
                        active_group = row
                    group_key = (name, active_group)
                previous_name = name

                if not registration:
                    unresolved = []
                    for key in _PARCEL_FIELDS:
                        if not direct.get(key):
                            continue
                        targets = merged_targets(row, key)
                        same_farm = name and targets and all(read(target, "name") == name
                                                            for target in targets)
                        if not same_farm or (key in ("area", "previous_registration", "lot_description")
                                             and len(targets) != 1):
                            unresolved.append(_LABELS[key])
                    if unresolved:
                        errors.append(f"Linha {row} possui {', '.join(unresolved)} mas nenhuma "
                                      "matrícula pôde ser determinada com segurança.")
                    if name and merged_targets(row, "name"):
                        property_rows.append(row)
                    elif any(direct.values()) and not unresolved:
                        warnings.append(f"Linha {row}: conteúdo sem campos estruturais suficientes "
                                        "ignorado; nenhuma entidade criada.")
                    continue

                anchor = registration_anchor(row)
                if not re.fullmatch(r"[A-Za-z0-9]+(?:[./-][A-Za-z0-9]+)*", registration) or registration == "0":
                    errors.append(f"Linha {anchor}: matrícula não identificável; revise o valor da célula.")
                    continue
                if not name:
                    errors.append(f"Linha {row}: matrícula sem fazenda vinculada; configure o agrupamento.")
                    continue
                record = records.setdefault(anchor, _ParcelRecord(anchor, name, registration, group_key))
                record.last_row = row
                if record.group_key != group_key:
                    errors.append(f"Linha {row}: matrícula mesclada atravessa grupos de fazenda.")
                    continue
                for key in (*_PARCEL_FIELDS, "municipality", "state"):
                    incoming = read(row, key)
                    if not incoming:
                        continue
                    if key in ("area", "previous_registration", "lot_description") and len(merged_targets(row, key)) > 1:
                        errors.append(f"Linha {row}: {_LABELS[key]} mesclada cobre várias matrículas; "
                                      "configure sua interpretação.")
                        continue
                    if key == "area":
                        area_column = columns[key]
                        area_cell = sheet.cell(*merges.get((row, area_column), (row, area_column)))
                        try:
                            incoming = format(parse_area_ha(area_cell.value), "f")
                        except ValueError as error:
                            errors.append(f"Linha {row}: {error}")
                            continue
                    previous = record.values.get(key)
                    if previous and previous != incoming:
                        errors.append(f"Linha {row}: {_LABELS[key]} conflitante no complemento da matrícula "
                                      f"iniciada na linha {anchor}.")
                    else:
                        record.values[key] = incoming
                observation = owner_observation(row)
                if observation.name or observation.document:
                    record.owners.add((observation.name, observation.document))
                    record.owner_observations.append(observation)

            # Only a real name merge may connect property-level rows to parcels.
            for row in property_rows:
                targets = merged_targets(row, "name")
                for target in targets:
                    if target not in records:
                        continue
                    record = records[target]
                    for key in ("municipality", "state"):
                        incoming = read(row, key, False)
                        if incoming:
                            existing = record.values.get(key)
                            if existing and incoming != existing:
                                errors.append(f"Linha {row}: {_LABELS[key]} conflitante no grupo da fazenda.")
                            else:
                                record.values[key] = incoming
                    owner = (read(row, "owner_name", False), read(row, "owner_document", False))
                    if any(owner):
                        record.owners.add(owner)
                        owner_targets = merged_targets(row, "owner_name") | merged_targets(row, "owner_document")
                        if target in owner_targets:
                            record.owner_observations.append(owner_observation(row))

            # A name-only continuation of the same physical merged name cell may
            # use its sole documented identity. Separate equal names never join.
            documents_by_name_cell: dict[str, set[str]] = {}
            for record in records.values():
                for observation in record.owner_observations:
                    document = normalized_owner_document(observation.document)
                    if observation.name and document:
                        documents_by_name_cell.setdefault(observation.name_cell, set()).add(document)
            owner_records: dict[str, PropertyOwner] = {}

            # Keep contiguous physical farm groups together, including optional
            # owner gaps. Never split a real name merge by changing owners.
            groups: list[list[_ParcelRecord]] = []
            for record in records.values():
                previous = groups[-1][-1] if groups else None
                shared_identity_merge = bool(previous and
                    previous.last_row + 1 == record.row and previous.name == record.name and
                    any(
                        read(record.row, key) and
                        {previous.row, record.row}.issubset(merged_targets(record.row, key)) and
                        all(read(target, "name") == record.name
                            for target in merged_targets(record.row, key))
                        for key in ("ccir", "itr", "car")
                    ))
                # Adjacent registration blocks may have separate farm-name
                # merges. A real shared document merge establishes their link;
                # matching names or document strings alone never do.
                if not previous or (previous.group_key != record.group_key and
                                    not shared_identity_merge):
                    groups.append([])
                groups[-1].append(record)
            properties: dict[str, RuralProperty] = {}
            parcel_ids: set[str] = set()
            for group in groups:
                owners = _legacy_identity_owners(set().union(*(record.owners for record in group)))
                # Retain the old identity inputs for previously valid single-owner
                # groups; multiple owners never select a first/principal person.
                owner_name, owner_document = next(iter(owners)) if len(owners) == 1 else ("", "")
                details = {}
                for key in ("municipality", "state", "ccir", "itr", "car"):
                    incoming = {record.values[key] for record in group if record.values.get(key)}
                    if len(incoming) > 1 and key in ("municipality", "state"):
                        errors.append(f"Linha {group[0].row}: {_LABELS[key]} conflitante no grupo de fazenda.")
                    if key in ("ccir", "itr", "car") and any(not record.values.get(key) for record in group):
                        incoming.add("")
                    details[key] = next(iter(incoming)) if len(incoming) == 1 else ""
                values = {
                    "source_file": relative_path, "profile": profile.name,
                    "owner_document": owner_document or owner_name,
                    "owner_name": owner_name, "name": group[0].name, **details,
                }
                property_id = _id("rp_", profile.property_identity_fields, values)
                if property_id in properties:
                    errors.append(f"Linha {group[0].row}: Blocos separados têm a mesma identidade de fazenda.")
                    continue
                property_item = RuralProperty(
                    external_id=property_id, name=group[0].name,
                    owner_name=owner_name, owner_document=owner_document,
                    source_file=relative_path, source_profile=profile.name, **details,
                    owners_normalized=True,
                )
                properties[property_id] = property_item
                for record in group:
                    parcel_values = {
                        "property_external_id": property_id,
                        "registration": record.registration,
                        **{key: record.values.get(key, "") for key in _PARCEL_FIELDS},
                    }
                    parcel_id = _id("pc_", profile.parcel_identity_fields, parcel_values)
                    if parcel_id in parcel_ids:
                        errors.append(f"Linha {record.row}: Matrículas com identidade duplicada; revise campos do perfil.")
                        continue
                    parcel_ids.add(parcel_id)
                    try:
                        area = parse_area_ha(record.values.get("area"))
                    except ValueError as error:
                        errors.append(f"Linha {record.row}: {error}")
                        continue
                    parcel = PropertyParcel(
                        external_id=parcel_id, property_external_id=property_id,
                        registration=record.registration,
                        previous_registration=record.values.get("previous_registration", ""),
                        area=area, lot_description=record.values.get("lot_description", ""),
                        extra_fields={key: record.values[key] for key in ("ccir", "itr", "car")
                                      if record.values.get(key)},
                    )
                    links: dict[str, ParcelOwner] = {}
                    for observation in record.owner_observations:
                        identity_document = observation.document
                        if not identity_document and observation.name:
                            documented = documents_by_name_cell.get(observation.name_cell, set())
                            if len(documented) == 1:
                                identity_document = next(iter(documented))
                            elif len(documented) > 1:
                                errors.append(f"Linha {record.row}: proprietário sem documento em célula "
                                              "compartilhada por documentos distintos; associação ambígua.")
                                continue
                        source_cell = observation.name_cell if observation.name else observation.document_cell
                        identifier = owner_id(identity_document, (
                            relative_path, sheet.title, source_cell,
                            observation.document_cell if observation.document else "",
                        ))
                        stored_owner = owner_records.setdefault(identifier, PropertyOwner(
                            identifier, observation.name, identity_document,
                        ))
                        if not stored_owner.name:
                            stored_owner.name = observation.name
                        property_item.owner_records[identifier] = owner_records[identifier]
                        link = links.setdefault(identifier, ParcelOwner(parcel_id, identifier))
                        if observation.name and observation.name not in link.source_names:
                            link.source_names += (observation.name,)
                        if observation.document and observation.document not in link.source_documents:
                            link.source_documents += (observation.document,)
                    parcel.owner_links = list(links.values())
                    property_item.parcels.append(parcel)
                if property_item.owner_count > 1:
                    property_item.owner_name = property_item.owner_document = ""
                elif property_item.owner_count == 1:
                    only_owner = property_item.owners[0]
                    property_item.owner_name = only_owner.name
                    property_item.owner_document = only_owner.document
            if errors:
                diagnostics = tuple(dict.fromkeys(errors))
                raise NeedsConfigurationError("\n".join((*diagnostics, *warnings)),
                                              diagnostics=diagnostics, warnings=tuple(warnings))
            if not properties:
                raise NeedsConfigurationError("Nenhuma fazenda com matrícula encontrada para o perfil.",
                                              warnings=tuple(warnings))
            only_owner = next(iter(owner_records.values())) if len(owner_records) == 1 else None
            return ParsedWorkbook(
                source_file=relative_path, profile_name=profile.name,
                owner_name=only_owner.name if only_owner else "",
                owner_document=only_owner.document if only_owner else "",
                properties=tuple(properties.values()), warnings=tuple(warnings),
            )
        finally:
            workbook.close()
