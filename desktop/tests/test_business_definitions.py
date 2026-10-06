"""Approved business rules, using only synthetic proposals and property data."""
from copy import copy
from decimal import Decimal
from hashlib import sha256
import json
import os
from pathlib import Path
import sqlite3
import subprocess
import sys

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.print_settings import PrintArea
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QDialog
from sqlalchemy.exc import IntegrityError

from amazon_agro.config.settings import AppSettings
from amazon_agro.domain.models import (
    ParticipantType, Proposal, ProposalProperty, ProposalPropertyParcel,
    PropertyClassification, PropertyParcel, RuralProperty,
)
from amazon_agro.exporters.excel import OpenpyxlExcelProposalExporter
from amazon_agro.exporters.pdf import SpreadsheetPdfProposalExporter
from amazon_agro.exporters.pdf_backends import ExcelComPdfBackend, PdfBackendSelector
from amazon_agro.repositories.fake_properties import FakePropertyRepository
from amazon_agro.repositories.sqlite_proposals import SQLiteProposalRepository
from amazon_agro.services.export_validator import ProposalExportValidator
from amazon_agro.services.proposal_service import ProposalService
from amazon_agro.ui.pages import ParticipantsPage
from amazon_agro.ui.property_selection import PropertiesPage, PropertyDetailsDialog


def proposal_with_parcels(count):
    proposal = Proposal(
        numero_proposta="TEST-PAGINATION", banco="Banco da Amazônia",
        proponente="Produtor Exemplo", cpf_cnpj="000.000.000-00",
        tecnico="Técnico Exemplo", finalidade="Custeio", cidade="Palmas",
        valor_total=Decimal("1000"), recursos_proprios=Decimal("100"),
        percentual_recursos_proprios=Decimal("10"),
    )
    proposal.properties = [ProposalProperty(
        proposal.id, "farm", None, property_name_snapshot="Fazenda Exemplo",
        municipality_snapshot="Palmas", state_snapshot="TO",
        selected_parcels=[ProposalPropertyParcel(
            f"parcel-{index}", f"MAT-{index:03}",
            classificacao=list(PropertyClassification)[(index - 1) % 3],
        ) for index in range(1, count + 1)],
    )]
    return proposal


def page_rows(count):
    return [(28, 29, 30, 31)] + [
        tuple(range(44 + (page - 1) * 16, 48 + (page - 1) * 16))
        for page in range(1, (count + 3) // 4)
    ]


def printed_page_areas(sheet):
    return len(PrintArea.from_string(sheet.print_area).ranges)


@pytest.fixture
def app(monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    application = QApplication.instance() or QApplication([])
    yield application
    application.processEvents()


def test_select_two_of_six_and_classify_each_parcel(app, tmp_path):
    farm = RuralProperty("farm", "Fazenda Exemplo", "Palmas", "TO", parcels=[
        PropertyParcel(f"parcel-{index}", "farm", f"MAT-{index}") for index in range(6)
    ])
    dialog = PropertyDetailsDialog(farm)
    assert dialog.selected_parcels() == []
    for row, code in ((1, "1"), (4, "3")):
        dialog.table.item(row, 0).setCheckState(Qt.CheckState.Checked)
        choice = dialog.table.cellWidget(row, 5)
        choice.setCurrentIndex(choice.findData(code))
    selected = dialog.selected_snapshots()
    assert [(parcel.parcel_external_id, parcel.registration_snapshot, parcel.classificacao)
            for parcel in selected] == [
        ("parcel-1", "MAT-1", PropertyClassification.CLASS_1),
        ("parcel-4", "MAT-4", PropertyClassification.FIDUCIARY_ALIENATION),
    ]
    proposal = proposal_with_parcels(0)
    proposal.properties[0].selected_parcels = selected
    output = OpenpyxlExcelProposalExporter(AppSettings(), FakePropertyRepository()).export(
        proposal, tmp_path / "two.xlsx")
    sheet = load_workbook(output).active
    assert [sheet[f"I{row}"].value for row in (28, 29, 30, 31)] == ["MAT-1", "MAT-4", None, None]
    assert [sheet[f"K{row}"].value for row in (28, 29)] == [1, 3]
    dialog.close()


def test_snapshot_reopens_selection_and_classifications_without_catalog(app, tmp_path, monkeypatch):
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    source = FakePropertyRepository([])
    service = ProposalService(repository, source)
    proposal = proposal_with_parcels(2)
    service.save(proposal)
    reopened = service.get(proposal.id)
    assert reopened.properties[0].selected_parcels == proposal.properties[0].selected_parcels
    page = PropertiesPage(service, AppSettings())
    page.load(reopened)
    assert page.collect(proposal.id)[0].selected_parcels == proposal.properties[0].selected_parcels
    def edit(dialog):
        assert [item.registration_snapshot for item in dialog.selected_snapshots()] == ["MAT-001", "MAT-002"]
        dialog.table.item(1, 0).setCheckState(Qt.CheckState.Unchecked)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(PropertyDetailsDialog, "exec", edit)
    page.edit_parcels(0)
    edited = page.collect(proposal.id)
    assert [item.registration_snapshot for item in edited[0].selected_parcels] == ["MAT-001"]
    assert len(service.get(proposal.id).properties[0].selected_parcels) == 2
    with repository.engine.connect() as connection:
        assert connection.exec_driver_sql("PRAGMA integrity_check").scalar() == "ok"
        assert connection.exec_driver_sql("PRAGMA foreign_key_check").fetchall() == []
    page.close()
    repository.engine.dispose()


def test_legacy_parcel_classification_migration_is_explicit_and_idempotent(tmp_path):
    path = tmp_path / "legacy.sqlite3"
    repository = SQLiteProposalRepository(path)
    proposal = proposal_with_parcels(2)
    proposal.properties[0].classificacao = PropertyClassification.FIDUCIARY_ALIENATION
    repository.save(proposal)
    with repository.engine.begin() as connection:
        connection.exec_driver_sql("ALTER TABLE proposal_property_parcels DROP COLUMN classificacao")
    repository.engine.dispose()
    migrated = SQLiteProposalRepository(path)
    reopened = migrated.get(proposal.id)
    assert [parcel.classificacao for parcel in reopened.properties[0].selected_parcels] == [
        PropertyClassification.FIDUCIARY_ALIENATION,
        PropertyClassification.FIDUCIARY_ALIENATION,
    ]
    reopened.properties[0].selected_parcels[1].classificacao = PropertyClassification.CREDIT_OBJECT
    migrated.save(reopened)
    migrated.engine.dispose()
    again = SQLiteProposalRepository(path)
    assert again.get(proposal.id).properties[0].selected_parcels[1].classificacao == PropertyClassification.CREDIT_OBJECT
    assert again.list_recent()[0].id == proposal.id
    again.engine.dispose()


def test_farm_default_change_does_not_reclassify_saved_parcels(tmp_path):
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    proposal = proposal_with_parcels(2)
    repository.save(proposal)
    reopened = repository.get(proposal.id)
    reopened.properties[0].classificacao = PropertyClassification.FIDUCIARY_ALIENATION
    repository.save(reopened)
    historical = repository.get(proposal.id)
    path = OpenpyxlExcelProposalExporter(AppSettings(), FakePropertyRepository([])).export(
        historical, tmp_path / "historical.xlsx")
    assert [load_workbook(path).active[f"K{row}"].value for row in (28, 29)] == [1, 2]
    repository.engine.dispose()


def test_failed_legacy_classification_migration_rolls_back_schema_and_data(tmp_path):
    path = tmp_path / "rollback.sqlite3"
    repository = SQLiteProposalRepository(path)
    proposal = proposal_with_parcels(1)
    repository.save(proposal)
    with repository.engine.begin() as connection:
        connection.exec_driver_sql("ALTER TABLE proposal_property_parcels DROP COLUMN classificacao")
        connection.exec_driver_sql(
            "CREATE TRIGGER fail_migration BEFORE UPDATE ON proposal_property_parcels "
            "BEGIN SELECT RAISE(ABORT, 'synthetic migration failure'); END"
        )
    repository.engine.dispose()
    with pytest.raises(IntegrityError, match="synthetic migration failure"):
        SQLiteProposalRepository(path)
    connection = sqlite3.connect(path)
    try:
        assert "classificacao" not in {row[1] for row in connection.execute(
            "PRAGMA table_info(proposal_property_parcels)")}
        assert connection.execute("SELECT COUNT(*) FROM proposals").fetchone() == (1,)
        assert connection.execute("SELECT registration_snapshot FROM proposal_property_parcels").fetchone() == ("MAT-001",)
        connection.execute("DROP TRIGGER fail_migration")
        connection.commit()
    finally:
        connection.close()
    recovered = SQLiteProposalRepository(path)
    assert recovered.get(proposal.id).properties[0].selected_parcels[0].classificacao == PropertyClassification.CLASS_1
    recovered.engine.dispose()


def test_missing_individual_classification_is_not_replaced_by_farm_default():
    proposal = proposal_with_parcels(1)
    proposal.properties[0].classificacao = PropertyClassification.CLASS_1
    proposal.properties[0].selected_parcels[0].classificacao = None
    with pytest.raises(ValueError, match="Classificação de matrícula"):
        proposal.validate()


def test_official_hipoteca_migrates_old_labels_without_changing_other_codes(tmp_path):
    config = tmp_path / "config.json"
    config.write_text(json.dumps({
        "property_classifications": {"1": "Garantia/Hipoteca (confirmar)", "2": "Crédito", "3": "Alienação"},
        "property_classification_labels": {
            "xlsx": {"1": "GARANTIA", "2": "CRÉDITO", "3": "ALIENAÇÃO"},
            "docx": {"1": "HIPOTECA", "2": "CRÉDITO", "3": "ALIENAÇÃO"},
        },
    }), encoding="utf-8")
    settings = AppSettings.load(config)
    assert settings.property_classifications["1"] == "Hipoteca"
    assert settings.classification_labels("xlsx")["1"] == "HIPOTECA"
    assert settings.property_classifications["2"] == "Crédito"
    settings.save()
    assert AppSettings.load(config).classification_labels("xlsx")["1"] == "HIPOTECA"


@pytest.mark.parametrize("count, expected", [(1, 1), (4, 1), (5, 2), (8, 2), (9, 3), (13, 4)])
def test_xlsx_pages_preserve_rows_styles_and_own_funds_indicator(tmp_path, count, expected):
    settings = AppSettings()
    template_hash = sha256(settings.template_path().read_bytes()).hexdigest()
    path = OpenpyxlExcelProposalExporter(settings, FakePropertyRepository()).export(
        proposal_with_parcels(count), tmp_path / f"parcels-{count}.xlsx")
    sheet = load_workbook(path).active
    assert printed_page_areas(sheet) == expected
    assert [brk.id for brk in sheet.row_breaks.brk] == [35 + 16 * index for index in range(expected - 1)]
    assert sheet.page_setup.fitToHeight == (1 if expected == 1 else 0)
    rows = [row for page in page_rows(count) for row in page]
    assert [sheet[f"I{row}"].value for row in rows[:count]] == [f"MAT-{i:03}" for i in range(1, count + 1)]
    assert [sheet[f"K{row}"].value for row in rows[:count]] == [(i - 1) % 3 + 1 for i in range(1, count + 1)]
    assert all(sheet[f"I{row}"].value is None for row in rows[count:])
    for page in page_rows(count)[1:]:
        row = page[0]
        for column in ("A", "E", "I", "K"):
            source, target = sheet[f"{column}28"], sheet[f"{column}{row}"]
            assert copy(source.font) == copy(target.font)
            assert copy(source.border) == copy(target.border)
            assert copy(source.alignment) == copy(target.alignment)
        assert f"A{row}:D{row}" in {str(item) for item in sheet.merged_cells.ranges}
        assert sheet.row_dimensions[row].height == sheet.row_dimensions[28].height
        assert sheet[f"A{row - 2}"].value == "IV - IMÓVEIS (CONTINUAÇÃO)"
        assert sheet[f"I{row - 1}"].value == sheet["I27"].value
    assert sheet["G21"].value == "Sim"
    assert sheet["H21"].value == pytest.approx(0.10)
    assert sheet["A32"].value == "1 - HIPOTECA"
    assert "R$ 100,00" not in {str(cell.value) for row in sheet for cell in row}
    assert sha256(settings.template_path().read_bytes()).hexdigest() == template_hash


@pytest.mark.parametrize("count", [1, 4, 5, 8, 9])
def test_pdf_conversion_keeps_xlsx_page_boundaries(tmp_path, count):
    path = OpenpyxlExcelProposalExporter(AppSettings(), FakePropertyRepository()).export(
        proposal_with_parcels(count), tmp_path / "source.xlsx")
    original_bytes = path.read_bytes()
    original = load_workbook(path).active
    class InspectBackend:
        name = "Test converter"
        def is_available(self): return True
        def convert(self, source, target):
            ready = load_workbook(source).active
            assert ready.print_area == original.print_area
            assert ready.row_breaks == original.row_breaks
            assert ready.page_setup == original.page_setup
            target.write_bytes(b"%PDF-synthetic")
    SpreadsheetPdfProposalExporter(PdfBackendSelector([InspectBackend()])).export(path, tmp_path / "output.pdf")
    assert path.read_bytes() == original_bytes


def test_amazon_default_is_user_classified_editable_and_removable(app, tmp_path):
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    service = ProposalService(repository, FakePropertyRepository())
    settings = AppSettings()
    proposal = service.new_proposal(settings)
    assert len(proposal.participants) == 1
    person = proposal.participants[0]
    assert person.proposal_id == proposal.id
    assert person.nome == "Amazon Agro Consultoria e Projetos LTDA"
    assert person.cpf_cnpj == "" and person.tipo == ParticipantType.MAIN_ISSUER
    assert "Tipo de participante inválido." not in ProposalExportValidator(settings, FakePropertyRepository()).validate(proposal).errors
    page = ParticipantsPage()
    page.load(proposal)
    page.table.item(0, 0).setText("Consultoria Exemplo")
    choice = page.table.cellWidget(0, 2)
    assert choice.currentData() == int(ParticipantType.MAIN_ISSUER)
    assert choice.isEnabled()
    choice.setCurrentIndex(choice.findData(int(ParticipantType.CUSTODIAN)))
    proposal.participants = page.collect(proposal.id)
    service.save(proposal)
    assert service.get(proposal.id).participants[0].nome == "Consultoria Exemplo"
    assert service.get(proposal.id).participants[0].tipo == ParticipantType.CUSTODIAN
    page.table.selectRow(0)
    page.remove_selected()
    proposal.participants = page.collect(proposal.id)
    service.save(proposal)
    for _ in range(3):
        proposal = service.get(proposal.id)
        page.load(proposal)
        proposal.participants = page.collect(proposal.id)
        assert proposal.participants == []
        service.save(proposal)
    assert service.new_proposal(settings).participants[0].id != person.id
    page.close()
    repository.engine.dispose()


def test_participant_defaults_accept_later_official_data(tmp_path):
    settings = AppSettings(
        default_consultancy_name="Nome configurado", default_consultancy_document="DOCUMENTO FORNECIDO",
        default_consultancy_type=int(ParticipantType.CUSTODIAN),
    )
    settings.save(tmp_path / "config.json")
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    proposal = ProposalService(repository, FakePropertyRepository()).new_proposal(AppSettings.load(tmp_path / "config.json"))
    assert proposal.participants[0].nome == "Nome configurado"
    assert proposal.participants[0].cpf_cnpj == "DOCUMENTO FORNECIDO"
    assert proposal.participants[0].tipo == ParticipantType.CUSTODIAN
    repository.engine.dispose()


def test_older_config_without_consultancy_type_uses_provisional_code_one(tmp_path):
    config = tmp_path / "config.json"
    config.write_text(json.dumps({"default_consultancy_document": ""}), encoding="utf-8")
    settings = AppSettings.load(config)
    assert settings.default_consultancy_type == int(ParticipantType.MAIN_ISSUER)
    settings.save()
    assert json.loads(config.read_text(encoding="utf-8"))["default_consultancy_type"] == 1
    assert AppSettings.load(config).default_consultancy_document == ""


@pytest.mark.parametrize("code", [0, 9, None, "1"])
def test_config_rejects_invalid_consultancy_default_type(tmp_path, code):
    config = tmp_path / "config.json"
    config.write_text(json.dumps({"default_consultancy_type": code}), encoding="utf-8")
    with pytest.raises(ValueError, match="tipo padrão de participante válido"):
        AppSettings.load(config)


def test_new_proposals_save_and_export_without_manual_participant_type_selection(app, tmp_path):
    settings = AppSettings()
    source = FakePropertyRepository([])
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    service = ProposalService(repository, source)
    validator = ProposalExportValidator(settings, source)
    exporter = OpenpyxlExcelProposalExporter(settings, source)
    page = ParticipantsPage()
    participant_ids = set()
    for index in range(3):
        proposal = service.new_proposal(settings)
        page.load(proposal)
        assert page.table.cellWidget(0, 2).currentData() == 1
        proposal.participants = page.collect(proposal.id)
        proposal.numero_proposta = f"TEST-DEFAULT-{index}"
        proposal.proponente = "Produtor Exemplo"
        proposal.cpf_cnpj = "000.000.000-00"
        proposal.tecnico = "Técnico Exemplo"
        proposal.finalidade = "Custeio"
        proposal.valor_total = Decimal("1000")
        assert validator.validate(proposal).ok
        service.save(proposal)
        reopened = service.get(proposal.id)
        assert reopened.participants[0].tipo == ParticipantType.MAIN_ISSUER
        participant_ids.add(reopened.participants[0].id)
        workbook = load_workbook(exporter.export(reopened, tmp_path / f"default-{index}.xlsx"))
        assert workbook.active["A9"].value == settings.default_consultancy_name
        assert workbook.active["M9"].value == 1 and workbook.active["I9"].value is None
        workbook.close()
    assert len(participant_ids) == 3
    page.close()
    repository.engine.dispose()


def test_reopened_proposal_keeps_one_amazon_and_its_saved_type(app, tmp_path):
    settings = AppSettings()
    repository = SQLiteProposalRepository(tmp_path / "proposals.sqlite3")
    service = ProposalService(repository, FakePropertyRepository([]))
    proposal = service.new_proposal(settings)
    participant_id = proposal.participants[0].id
    service.save(proposal)
    settings.default_consultancy_type = int(ParticipantType.CUSTODIAN)
    assert service.new_proposal(settings).participants[0].tipo == ParticipantType.CUSTODIAN
    page = ParticipantsPage()
    for _ in range(3):
        reopened = service.get(proposal.id)
        page.load(reopened)
        reopened.participants = page.collect(reopened.id)
        assert len(reopened.participants) == 1
        assert reopened.participants[0].id == participant_id
        assert reopened.participants[0].tipo == ParticipantType.MAIN_ISSUER
        service.save(reopened)
    page.close()
    repository.engine.dispose()


def test_all_banks_use_identical_required_fields():
    validator = ProposalExportValidator(AppSettings(), FakePropertyRepository())
    results = []
    for bank in ("Banco da Amazônia", "Banco do Brasil", "Outro banco"):
        proposal = proposal_with_parcels(9)
        proposal.banco = bank
        assert validator.validate(proposal).ok
        proposal.tecnico = ""
        results.append(validator.validate(proposal).errors)
    assert results[0] == results[1] == results[2] == ("Informe o técnico responsável.",)


PDF_READER_PYTHON = os.environ.get("AMAZON_AGRO_PDF_READER_PYTHON")


@pytest.mark.skipif(not PDF_READER_PYTHON or not ExcelComPdfBackend().is_available(),
                    reason="Native Excel COM and a Python with pypdf are required")
@pytest.mark.parametrize("count, expected", [(1, 1), (4, 1), (5, 2), (8, 2), (9, 3)])
def test_excel_com_pdf_matches_xlsx_pages(tmp_path, count, expected):
    path = OpenpyxlExcelProposalExporter(AppSettings(), FakePropertyRepository()).export(
        proposal_with_parcels(count), tmp_path / f"native-{count}.xlsx")
    pdf = tmp_path / f"native-{count}.pdf"
    # Isolate real Office from the suite's mocked COM modules and Qt lifetime.
    subprocess.run([
        sys.executable, "-c",
        "from pathlib import Path; import sys; "
        "from amazon_agro.exporters.pdf import SpreadsheetPdfProposalExporter; "
        "from amazon_agro.exporters.pdf_backends import PdfBackendSelector, ExcelComPdfBackend; "
        "SpreadsheetPdfProposalExporter(PdfBackendSelector([ExcelComPdfBackend()])).export(Path(sys.argv[1]), Path(sys.argv[2]))",
        str(path), str(pdf),
    ], capture_output=True, text=True, check=True, timeout=120)
    result = subprocess.run([
        PDF_READER_PYTHON, "-c",
        "import json,sys; from pypdf import PdfReader; print(json.dumps([p.extract_text() for p in PdfReader(sys.argv[1]).pages]))",
        str(pdf),
    ], capture_output=True, text=True, check=True, timeout=30)
    pages = json.loads(result.stdout)
    assert len(pages) == expected == printed_page_areas(load_workbook(path).active)
    for number, text in enumerate(pages):
        for index in range(number * 4 + 1, min(count, number * 4 + 4) + 1):
            assert f"MAT-{index:03}" in text
        assert "HIPOTECA" in text
        if number:
            assert "CONTINUAÇÃO" in text
