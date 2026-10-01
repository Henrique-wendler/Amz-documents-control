"""Read-only workbook inspection for mapping and future configuration UI."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from amazon_agro.integrations.workbook_profile import PropertyWorkbookProfile, normalized


class NeedsConfigurationError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class HeaderCandidate:
    sheet: str
    row: int
    columns: dict[str, int]
    duplicates: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class SheetInspection:
    name: str
    first_rows: tuple[tuple[str, ...], ...]
    header_candidates: tuple[HeaderCandidate, ...]
    merged_ranges: tuple[str, ...]
    column_names: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class WorkbookInspection:
    path: Path
    sheets: tuple[SheetInspection, ...]


def header_candidates(sheet: Worksheet, profile: PropertyWorkbookProfile) -> list[HeaderCandidate]:
    candidates: list[HeaderCandidate] = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, profile.header_scan_rows)):
        columns: dict[str, int] = {}
        duplicates: set[str] = set()
        for cell in row:
            field_name = profile.field_for_header(cell.value)
            if field_name:
                if field_name in columns:
                    duplicates.add(field_name)
                columns[field_name] = cell.column
        if len(columns) >= 2:
            candidates.append(HeaderCandidate(
                sheet.title, row[0].row, columns, tuple(sorted(duplicates))
            ))
    return sorted(candidates, key=lambda item: (-len(item.columns), item.row))


def select_header(workbook: object, profile: PropertyWorkbookProfile) -> tuple[Worksheet, HeaderCandidate]:
    sheets = workbook.worksheets
    if profile.sheet_selector:
        sheets = [
            sheet for sheet in sheets
            if normalized(sheet.title) == normalized(profile.sheet_selector)
        ]
    required = set(profile.required_fields)
    if profile.owner_name_cell:
        required.discard("owner_name")
    matches = [
        (sheet, candidate)
        for sheet in sheets for candidate in header_candidates(sheet, profile)
        if required <= set(candidate.columns)
    ]
    if matches:
        if not profile.sheet_selector and len({sheet.title for sheet, _ in matches}) > 1:
            raise NeedsConfigurationError(
                "Mais de uma aba contém imóveis; selecione a aba no perfil."
            )
        best_size = max(len(candidate.columns) for _, candidate in matches)
        best = [(sheet, candidate) for sheet, candidate in matches
                if len(candidate.columns) == best_size]
        if len(best) == 1 and not best[0][1].duplicates:
            return best[0]
        raise NeedsConfigurationError(
            "Mais de um cabeçalho possível ou colunas duplicadas; selecione a aba/perfil."
        )
    raise NeedsConfigurationError(
        "Cabeçalhos obrigatórios ausentes, duplicados ou aba incompatível com o perfil."
    )


def _preview(value: object, field_name: str = "") -> str:
    text = str(value or "")
    if field_name == "owner_name":
        parts = text.split()
        return (parts[0] + " " + parts[1][:1] + ".") if len(parts) > 1 else text[:1] + "."
    digits = re.sub(r"\D", "", text)
    if field_name == "owner_document" or (len(digits) in {11, 14} and len(text) <= 22):
        if not digits:
            return "***"
        return digits[:3] + "***" + digits[-2:]
    return text[:100]


class WorkbookInspector:
    def inspect(
        self, path: Path, profile: PropertyWorkbookProfile,
        preview_rows: int = 12,
    ) -> WorkbookInspection:
        workbook = load_workbook(path, data_only=True, read_only=False)
        try:
            result: list[SheetInspection] = []
            for sheet in workbook.worksheets:
                candidates = header_candidates(sheet, profile)
                best = candidates[0] if candidates else None
                sensitive_columns = {
                    column: field_name
                    for field_name, column in (best.columns.items() if best else ())
                    if field_name in {"owner_name", "owner_document"}
                }
                first_rows = tuple(
                    tuple(
                        _preview(
                            cell.value,
                            sensitive_columns.get(cell.column, "")
                            if best and cell.row > best.row else "",
                        )
                        for cell in row[:20]
                    )
                    for row in sheet.iter_rows(max_row=min(preview_rows, sheet.max_row))
                )
                names = tuple(
                    _preview(sheet.cell(best.row, col).value)
                    for col in range(1, sheet.max_column + 1)
                ) if best else ()
                result.append(SheetInspection(
                    name=sheet.title, first_rows=first_rows,
                    header_candidates=tuple(candidates),
                    merged_ranges=tuple(str(item) for item in sheet.merged_cells.ranges),
                    column_names=names,
                ))
            return WorkbookInspection(Path(path), tuple(result))
        finally:
            workbook.close()
